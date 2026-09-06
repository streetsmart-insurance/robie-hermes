import openpyxl
import json

wb = openpyxl.load_workbook("/Users/carloferrara/Documents/antigravity/happy-fermi/data/APE_policies.xlsx", data_only=False)

print("Sheet names:", wb.sheetnames)

all_data = {}

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n=================== SHEET: {sheet_name} (max_row={ws.max_row}, max_column={ws.max_column}) ===================")
    rows_data = []
    for r in range(1, ws.max_row + 1):
        row_vals = []
        has_content = False
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            val = cell.value
            hyperlink = cell.hyperlink.target if cell.hyperlink else None
            
            if val is not None or hyperlink is not None:
                has_content = True
                val_str = str(val).strip() if val is not None else ""
                if hyperlink:
                    row_vals.append(f"{val_str} [LINK: {hyperlink}]")
                else:
                    row_vals.append(val_str)
            else:
                row_vals.append("")
        
        # trim trailing empty cells
        while row_vals and not row_vals[-1]:
            row_vals.pop()
            
        if any(row_vals):
            rows_data.append(row_vals)
            print(f"R{r}: {' | '.join(row_vals)}")
            
    all_data[sheet_name] = rows_data

with open("/Users/carloferrara/Documents/antigravity/happy-fermi/data/APE_parsed_data.json", "w") as f:
    json.dump(all_data, f, indent=2)

print("\nSaved parsed data to data/APE_parsed_data.json")
