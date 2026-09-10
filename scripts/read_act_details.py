import openpyxl

wb = openpyxl.load_workbook('/opt/renewal-automation-system/data/input_reports/EZLynx_Scheduled_1a085ad6_Robie_-_EZLynx_Activities_2026-09-09T0605.xlsx', read_only=True)
ws = wb.active

target_apps = {'72885007', '21586424', '199729236', '38142343', '186530244', '27904633', '41600472'}

headers = [str(c) for c in next(ws.iter_rows(values_only=True))]
print('Columns:', list(enumerate(headers)))

for r in ws.iter_rows(values_only=True):
    r_str = [str(c) if c is not None else '' for c in r]
    row_text = ' '.join(r_str)
    for app in target_apps:
        if app in row_text:
            date_col = ''
            for c in r_str:
                if '2026-09' in c or '2026-08-3' in c:
                    date_col = c
                    break
            if date_col:
                print(f"\n==================================================")
                print(f"App: {app} | Date: {date_col}")
                for idx, val in enumerate(r_str):
                    if len(val.strip()) > 0 and idx < len(headers):
                        h = headers[idx]
                        if h not in ['Applicant Id', 'Applicant Name', 'Agency']:
                            print(f"  {h}: {val[:300]}")
