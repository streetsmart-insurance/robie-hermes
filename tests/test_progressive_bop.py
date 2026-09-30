"""Progressive BOP pending-cancel pull. No live FAO session."""
from __future__ import annotations

import io
import json
import os
import re
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from robie_job_engine.intake_core import IntakeHold
from robie_job_engine.progressive_agent_context import assert_agent_context as shared_assert_agent_context
from robie_job_engine.progressive_bop import (
    BOP_SCOPE,
    DEFAULT_OUTPUT_ROOT,
    DRIVE_QA_PARENT_ID,
    DRIVE_BOP_CHILD_NAME,
    MANAGE_POLICIES_CSS,
    PendingCancelPolicy,
    PolicyDocument,
    ReportCapture,
    RowHold,
    agent_codes_in_text,
    assert_agent_context,
    assess_bop_gate,
    build_parser,
    main,
    module_import_names,
    _allowed_pdf_url,
    _BopFrameSurface,
    navigate_to_pending_cancel,
    noc_filename,
    parse_excel_report,
    parse_pdf_report,
    parse_report_rows,
    parse_xls_export,
    apply_bop_report_dates,
    _PENDING_CANCEL_PDF_EXPORT,
    _PENDING_CANCEL_XLS_EXPORT,
    policies_from_report_text,
    read_report_from_page,
    report_from_extracted,
    require_loopback_cdp,
    select_fao_page,
    select_notice,
)
from robie_job_engine.progressive_fao_memo import assert_agent_context as fao_assert_agent_context


PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 16
DAY = date(2026, 9, 26)
PDF_A = b"%PDF-1.4\nnotice-a\n%%EOF\n"
PDF_B = b"%PDF-1.4\nnotice-b\n%%EOF\n"


def _excel(rows: list[list[object]]) -> bytes:
    book = Workbook()
    sheet = book.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _text_pdf(text: str) -> bytes:
    safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 700 Td ({safe}) Tj ET".encode("latin-1")
    chunks = [b"%PDF-1.4\n"]

    def obj(number: int, content: bytes) -> None:
        chunks.append(f"{number} 0 obj\n".encode() + content + b"\nendobj\n")

    obj(1, b"<< /Type /Catalog /Pages 2 0 R >>")
    obj(2, b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>")
    obj(3, b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>")
    obj(4, b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    obj(5, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    xref_pos = sum(len(chunk) for chunk in chunks)
    cursor = len(chunks[0])
    offsets = []
    for chunk in chunks[1:]:
        offsets.append(cursor)
        cursor += len(chunk)
    xref = [b"xref\n0 6\n", b"0000000000 65535 f \n"]
    xref.extend(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    trailer = f"trailer << /Size 6 /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode()
    return b"".join(chunks) + b"".join(xref) + trailer


class FakeLocator:
    def __init__(self, page, selector="", role=None, name=None):
        self.page = page
        self.selector = selector
        self.role = role
        self.name = name

    def count(self):
        if self.selector == "input[type='password']":
            return self.page.password_count
        if self.selector == "table":
            return 0
        if self.selector.startswith("embed"):
            return 0
        if self.role is not None:
            if isinstance(self.name, re.Pattern):
                return sum(
                    1
                    for role, name in self.page.roles
                    if role == self.role and self.name.search(str(name))
                )
            return 1 if (self.role, self.name) in self.page.roles else 0
        return 0

    def click(self):
        self.page.clicked.append((self.role, self.name))

    def inner_text(self):
        if self.selector == "body":
            return self.page.body
        return ""

    def nth(self, index):
        return self

    def all(self):
        return []

    def locator(self, selector):
        return FakeLocator(self.page, selector)

    def get_attribute(self, name):
        return None


class FakePage:
    def __init__(self, roles, *, body="CA33617", url="https://www.foragentsonly.com/home", popup=None):
        self.roles = set(roles)
        self.body = body
        self.url = url
        self.password_count = 0
        self.clicked = []
        self.popup = popup
        self.closed = False
        self.close_raises = False
        self.frames = None
        self.context = None

    def close(self):
        self.closed = True
        if self.close_raises:
            raise RuntimeError("HPLanding close failed")

    def goto(self, url, **kwargs):
        self.clicked.append(("goto", url))
        self.url = url

    def locator(self, selector):
        return FakeLocator(self, selector)

    def get_by_role(self, role, name=None, exact=True):
        return FakeLocator(self, role=role, name=name)

    def expect_popup(self, timeout=None):
        page = self

        class _Popup:
            def __enter__(self):
                return self

            def __exit__(self, exc_type, exc, tb):
                self.value = page.popup
                return False

        return _Popup()


class FakeBrowser:
    def __init__(self, capture, notices=None, root: Path | None = None):
        self.capture = capture
        self.notices = dict(notices or {})
        self.root = root
        self.downloaded: list[str] = []
        self.shell_blocked = None

    def load_report(self, report_date):
        if isinstance(self.capture, Exception):
            raise self.capture
        if self.root is not None:
            folder = self.root / report_date.isoformat()
            folder.mkdir(mode=0o700, exist_ok=True)
            os.chmod(folder, 0o700)
            (folder / noc_filename("860521214")).write_bytes(b"%PDF-1.4\nconflict\n")
        return self.capture

    def download_notice(self, policy):
        self.downloaded.append(policy.policy_number)
        item = self.notices[policy.policy_number]
        if isinstance(item, Exception):
            raise item
        return item


def _policy(number: str, insured: str = "ACME LLC") -> PendingCancelPolicy:
    return PendingCancelPolicy(number, insured, DAY)


def _report(policies, *, source="table", blank=False, observed=None):
    from robie_job_engine.progressive_bop import PendingCancelReport
    return PendingCancelReport(DAY, observed, tuple(policies), source, blank)


class ProgressiveBopTests(unittest.TestCase):
    def setUp(self):
        self._env = patch.dict(os.environ, {"ROBIE_ENV": "TEST"})
        self._env.start()
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name) / "qa"
        self.addCleanup(self._env.stop)
        self.addCleanup(self.tmp.cleanup)

    def test_scope_and_filename(self):
        self.assertEqual(BOP_SCOPE, "bop_pending_cancel_nonpayment")
        self.assertEqual(noc_filename("860521214"), "860521214 - NOC - Non Payment.pdf")
        self.assertEqual(DRIVE_QA_PARENT_ID, "1cLEpR-0T6KdiVjcdAr0qpGTO447MetI2")
        self.assertEqual(DRIVE_BOP_CHILD_NAME, "Progressive BOP")
        with self.assertRaises(IntakeHold):
            noc_filename("12")
        with self.assertRaises(IntakeHold):
            noc_filename("../860521214")

    def test_output_root_default_is_the_hermes_test_pack(self):
        args = build_parser().parse_args(["--report-date", "2026-09-26"])
        self.assertEqual(
            Path(args.output_root),
            DEFAULT_OUTPUT_ROOT,
        )
        self.assertTrue(str(DEFAULT_OUTPUT_ROOT).endswith(
            "carrier-pull-qa/progressive-bop"
        ))

    def test_module_does_not_import_ezlynx_or_drive_client(self):
        names = module_import_names()
        self.assertFalse(any("ezlynx" in name for name in names))
        self.assertFalse(any("googleapiclient" in name or name.startswith("google") for name in names))

    def test_table_policies_and_blank_header(self):
        report = parse_report_rows(
            [
                ("Report Date", "09/26/2026"),
                ("Policy Number", "Named Insured"),
                ("860521214", "3JR Contracting LLC"),
                (860521215, "ALTI TRANSPORT LLC"),
            ],
            report_date=DAY,
            source="table",
        )
        self.assertEqual([row.policy_number for row in report.policies], ["860521214", "860521215"])
        self.assertFalse(report.blank)
        blank = parse_report_rows(
            [("Policy Number", "Named Insured")],
            report_date=DAY,
            source="table",
        )
        self.assertTrue(blank.blank)
        self.assertEqual(blank.policies, ())

    def test_duplicate_policy_wrong_date_and_unparsed_text_hold(self):
        with self.assertRaises(IntakeHold):
            parse_report_rows(
                [("Policy Number",), ("860521214",), ("860521214",)],
                report_date=DAY,
                source="table",
            )
        with self.assertRaises(IntakeHold):
            parse_report_rows(
                [("Policy Number", "Cancel Date"), ("860521214", "09/25/2026")],
                report_date=DAY,
                source="table",
            )
        with self.assertRaises(IntakeHold):
            policies_from_report_text("Pending Cancel for Nonpayment\nPage 1", report_date=DAY)
        blank = policies_from_report_text("No records found", report_date=DAY)
        self.assertTrue(blank.blank)

    def test_excel_and_pdf_reports(self):
        excel = parse_excel_report(
            _excel([
                ["Policy Number", "Named Insured", "Cancel Date"],
                [860521214, "3JR Contracting LLC", date(2026, 9, 26)],
            ]),
            report_date=DAY,
        )
        self.assertEqual(excel.policies[0].filename, "860521214 - NOC - Non Payment.pdf")
        self.assertEqual(excel.source, "excel")
        blank = parse_excel_report(
            _excel([["Policy Number", "Named Insured"]]),
            report_date=DAY,
        )
        self.assertTrue(blank.blank)
        with self.assertRaises(IntakeHold):
            parse_excel_report(_excel([]), report_date=DAY)
        pdf = parse_pdf_report(
            _text_pdf("Policy Number 879512352 ALTI TRANSPORT LLC"),
            report_date=DAY,
        )
        self.assertEqual(pdf.policies[0].policy_number, "879512352")
        self.assertIn("ALTI", pdf.policies[0].insured_name)

    def test_two_policy_tables_hold_and_blank_page_text_is_empty(self):
        with self.assertRaises(IntakeHold):
            report_from_extracted(
                tables=[
                    (("Policy Number",), (("111111111",),)),
                    (("Policy Number",), (("222222222",),)),
                ],
                page_text="",
                excel_bytes=None,
                pdf_bytes=None,
                report_date=DAY,
            )
        report = report_from_extracted(
            tables=[],
            page_text="Pending Cancel for Nonpayment. No policies found.",
            excel_bytes=None,
            pdf_bytes=None,
            report_date=DAY,
        )
        self.assertTrue(report.blank)
        self.assertEqual(report.source, "page")

    def test_notice_selection_skips_other_docs_and_holds_missing(self):
        chosen = select_notice(
            [
                PolicyDocument("Declarations Page", DAY, 0),
                PolicyDocument("Notice of Non Payment", DAY, 1),
                PolicyDocument("Notice of Non-Payment", date(2026, 9, 25), 2),
            ],
            DAY,
        )
        self.assertEqual(chosen.row_index, 1)
        named = select_notice(
            [PolicyDocument("Notice of Non Payment 09/26/2026", None, 0)],
            DAY,
        )
        self.assertEqual(named.row_index, 0)
        with self.assertRaises(RowHold) as missing:
            select_notice([PolicyDocument("Invoice", DAY, 0)], DAY)
        self.assertIn("missing", str(missing.exception))
        with self.assertRaises(RowHold):
            select_notice(
                [
                    PolicyDocument("Notice of Non Payment", DAY, 0),
                    PolicyDocument("Notice of Nonpayment", DAY, 1),
                ],
                DAY,
            )
        with self.assertRaises(RowHold) as undated:
            select_notice([PolicyDocument("Notice of Non Payment", None, 0)], DAY)
        self.assertIn("date is missing", str(undated.exception))

    def test_gate_empty_pulled_and_held(self):
        from robie_job_engine.progressive_bop import PolicyOutcome
        self.assertEqual(assess_bop_gate(policy_count=0, outcomes=[], pdf_count=0, screenshot_ok=True), ("EMPTY", None))
        pulled = PolicyOutcome("860521214", "A", "id", "860521214 - NOC - Non Payment.pdf", "pulled")
        self.assertEqual(
            assess_bop_gate(policy_count=1, outcomes=[pulled], pdf_count=1, screenshot_ok=True)[0],
            "PULLED",
        )
        held = PolicyOutcome("860521214", "A", "id", "f.pdf", "held", reason="missing")
        status, reason = assess_bop_gate(policy_count=1, outcomes=[held], pdf_count=0, screenshot_ok=True)
        self.assertEqual(status, "HELD")
        self.assertIn("860521214", reason)
        mismatch = assess_bop_gate(policy_count=1, outcomes=[pulled], pdf_count=2, screenshot_ok=True)
        self.assertEqual(mismatch[0], "HELD")
        self.assertIn("does not match", mismatch[1])
        self.assertEqual(
            assess_bop_gate(policy_count=0, outcomes=[], pdf_count=0, screenshot_ok=False)[0],
            "HELD",
        )

    def test_loopback_and_single_fao_tab(self):
        self.assertEqual(require_loopback_cdp("http://127.0.0.1:9222"), "http://127.0.0.1:9222")
        with self.assertRaises(IntakeHold):
            require_loopback_cdp("http://10.0.0.8:9222")
        page = FakePage(set(), url="https://www.foragentsonly.com/home")
        self.assertIs(select_fao_page([page, FakePage(set(), url="https://example.test")]), page)
        with self.assertRaises(IntakeHold):
            select_fao_page([page, FakePage(set(), url="https://portal.foragentsonly.com/other")])
        with self.assertRaises(IntakeHold):
            select_fao_page([])

    def test_navigation_order_and_auth_hold(self):
        report = FakePage(
            {("link", "View Reports"), ("button", "Pending Cancel for Nonpayment")},
            url="https://www.foragentsonly.com/bop",
        )
        shell = FakePage(
            {("link", "Manage Policies"), ("link", "Businessowner/Contractor GL")},
            popup=report,
        )
        opened = navigate_to_pending_cancel(shell, "ca33617")
        self.assertIs(opened, report)
        self.assertEqual(
            shell.clicked,
            [("goto", "https://www.foragentsonly.com/landingpages/managepolicies/"), ("link", "Businessowner/Contractor GL")],
        )
        self.assertNotIn(("link", "Manage Policies"), shell.clicked)
        self.assertEqual(
            report.clicked,
            [("link", "View Reports"), ("button", "Pending Cancel for Nonpayment")],
        )
        locked = FakePage({("link", "Manage Policies")}, body="CA33617", url="https://foragentsonlylogin.progressive.com/login")
        with self.assertRaises(IntakeHold):
            navigate_to_pending_cancel(locked, "CA33617")
        self.assertEqual(locked.clicked, [])
        ambiguous = FakePage(
            {("link", "Manage Policies")},
            body="CA33617 CA11111",
        )
        with self.assertRaises(IntakeHold):
            navigate_to_pending_cancel(ambiguous, "CA33617")
        self.assertEqual(ambiguous.clicked, [])

    def test_bop_application_host_is_an_allowed_report_pdf_source(self):
        self.assertTrue(_allowed_pdf_url("https://bop.americanstrategic.com/report.pdf"))
        self.assertFalse(_allowed_pdf_url("http://bop.americanstrategic.com/report.pdf"))
        self.assertFalse(_allowed_pdf_url("https://evil.example/report.pdf"))

    def test_hplanding_popup_attaches_to_the_bop_app_and_closes(self):
        app = FakePage(
            {("button", "VIEW REPORTS"), ("link", "Pending Cancel for Nonpayment")},
            url="https://bop.americanstrategic.com/reports",
            body="VIEW REPORTS",
        )
        landing = FakePage(
            {("button", "Close this window")},
            url="https://www.foragentsonly.com/HPLanding",
            body="The application opened in another window.",
        )
        shell = FakePage(
            {("link", "Businessowner/Contractor GL")},
            url="https://www.foragentsonly.com/landingpages/managepolicies/",
            popup=landing,
        )
        shell.context = type("Ctx", (), {"pages": [shell, app]})()
        opened = navigate_to_pending_cancel(shell, "CA33617")
        self.assertIs(opened, app)
        self.assertTrue(landing.closed)
        self.assertEqual(
            app.clicked,
            [("button", "VIEW REPORTS"), ("link", "Pending Cancel for Nonpayment")],
        )
        self.assertNotIn(("link", "Manage Policies"), shell.clicked)
        self.assertEqual(landing.clicked, [])

    def test_hplanding_close_failure_still_attaches_to_the_bop_app(self):
        app = FakePage(
            {("link", "View Reports"), ("button", "Pending Cancel for Nonpayment")},
            url="https://bop.americanstrategic.com/",
        )
        landing = FakePage(
            set(),
            url="https://www.foragentsonly.com/home",
            body="Please Close this window",
        )
        landing.close_raises = True
        shell = FakePage(
            {("link", "Go to Businessowner/Contractor GL policy search")},
            url="https://www.foragentsonly.com/home",
            popup=landing,
        )
        shell.context = type("Ctx", (), {"pages": [app]})()
        opened = navigate_to_pending_cancel(shell, "CA33617")
        self.assertIs(opened, app)
        self.assertTrue(landing.closed)
        self.assertEqual(app.clicked[0], ("link", "View Reports"))

    def test_hplanding_frame_is_the_bop_application(self):
        frame = FakePage(
            {("link", "VIEW REPORTS"), ("button", "Pending Cancel for Nonpayment")},
            url="https://bop.americanstrategic.com/home",
        )
        landing = FakePage(
            {("button", "Close this window")},
            url="https://www.foragentsonly.com/HPLanding/session",
            body="Close this window",
        )
        landing.frames = [landing, frame]
        shell = FakePage(
            {("link", "Businessowner/Contractor GL")},
            url="https://www.foragentsonly.com/landingpages/managepolicies/home",
            popup=landing,
        )
        opened = navigate_to_pending_cancel(shell, "CA33617")
        self.assertIsInstance(opened, _BopFrameSurface)
        self.assertEqual(opened.url, frame.url)
        self.assertFalse(landing.closed)
        self.assertEqual(
            frame.clicked,
            [("link", "VIEW REPORTS"), ("button", "Pending Cancel for Nonpayment")],
        )

    def test_hplanding_waits_for_bop_americanstrategic_before_view_reports(self):
        app = FakePage(
            {("button", "VIEW REPORTS"), ("link", "Pending Cancel for Nonpayment")},
            url="https://bop.americanstrategic.com/",
        )
        landing = FakePage(
            {("button", "Close this window")},
            url="https://sbr1.foragentsonly.com/portal/HPLanding.aspx",
            body="Close this window",
        )
        shell = FakePage(
            {("link", "Businessowner/Contractor GL")},
            url="https://www.foragentsonly.com/landingpages/managepolicies/",
            popup=landing,
        )

        class _Ctx:
            def __init__(self):
                self.pages = [shell, landing]
                self.listeners = []

            def on(self, event, fn):
                if event == "page":
                    self.listeners.append(fn)

            def remove_listener(self, event, fn):
                if fn in self.listeners:
                    self.listeners.remove(fn)

        ctx = _Ctx()
        shell.context = ctx
        landing.context = ctx

        def wait(_ms):
            if app not in ctx.pages:
                ctx.pages.append(app)
                for fn in list(ctx.listeners):
                    fn(app)

        shell.wait_for_timeout = wait
        opened = navigate_to_pending_cancel(shell, "CA33617")
        self.assertIs(opened, app)
        self.assertTrue(landing.closed)
        self.assertEqual(landing.clicked, [])
        self.assertEqual(
            app.clicked,
            [("button", "VIEW REPORTS"), ("link", "Pending Cancel for Nonpayment")],
        )
        self.assertNotIn(("link", "Manage Policies"), shell.clicked)

    def test_hplanding_without_the_bop_app_holds_before_view_reports(self):
        landing = FakePage(
            {("button", "Close this window"), ("link", "View Reports")},
            url="https://www.foragentsonly.com/HPLanding",
            body="Close this window",
        )
        shell = FakePage(
            {("link", "Businessowner/Contractor GL")},
            url="https://www.foragentsonly.com/",
            popup=landing,
        )
        with self.assertRaisesRegex(IntakeHold, "HPLanding"):
            navigate_to_pending_cancel(shell, "CA33617")
        self.assertEqual(landing.clicked, [])

    def test_two_bop_windows_hold(self):
        first = FakePage(set(), url="https://bop.americanstrategic.com/one")
        second = FakePage(set(), url="https://bop.americanstrategic.com/two")
        landing = FakePage(
            {("button", "Close this window")},
            url="https://www.foragentsonly.com/HPLanding",
        )
        shell = FakePage(
            {("link", "Businessowner/Contractor GL")},
            url="https://www.foragentsonly.com/home",
            popup=landing,
        )
        shell.context = type("Ctx", (), {"pages": [first, second]})()
        with self.assertRaisesRegex(IntakeHold, "more than one BOP window"):
            navigate_to_pending_cancel(shell, "CA33617")
        self.assertFalse(landing.closed)

    def test_blank_page_text_does_not_require_an_export(self):
        page = FakePage(set(), body="Pending Cancel for Nonpayment. No records.")
        report = read_report_from_page(page, DAY)
        self.assertTrue(report.blank)
        self.assertEqual(report.policies, ())

    def _run(self, browser, *extra):
        stdout = io.StringIO()
        with patch("sys.stdout", stdout):
            code = main(
                ["--report-date", DAY.isoformat(), "--output-root", str(self.root), *extra],
                browser_factory=lambda args: browser,
            )
        payload = json.loads(stdout.getvalue() or "{}")
        return code, payload

    def test_blank_report_writes_empty_pack(self):
        browser = FakeBrowser(ReportCapture(PNG, _report((), blank=True, source="page"), None))
        code, payload = self._run(browser)
        self.assertEqual(code, 0)
        self.assertEqual(payload["status"], "EMPTY")
        self.assertEqual(payload["pdfs"], 0)
        self.assertEqual(payload["ezlynx"], "not_run")
        folder = self.root / DAY.isoformat()
        self.assertTrue((folder / "pending-cancel-nonpayment-2026-09-26.png").read_bytes().startswith(b"\x89PNG"))
        self.assertFalse(list(folder.glob("*.pdf")))
        manifest = json.loads((folder / "manifest.json").read_text())
        self.assertEqual(manifest["status"], "EMPTY")
        self.assertEqual(manifest["policies_on_report"], 0)
        self.assertEqual(manifest["drive"]["parent_id"], DRIVE_QA_PARENT_ID)
        self.assertEqual(manifest["drive"]["child_name"], "Progressive BOP")
        self.assertEqual(manifest["drive"]["status"], "not_run")
        self.assertIn("No policies", (folder / "README.md").read_text())
        self.assertEqual(browser.downloaded, [])

    def test_pulls_matching_notices_and_holds_the_missing_row(self):
        policies = (_policy("860521214", "3JR Contracting LLC"), _policy("879512352", "ALTI TRANSPORT LLC"))
        browser = FakeBrowser(
            ReportCapture(PNG, _report(policies), None),
            {
                "860521214": PDF_A,
                "879512352": RowHold("Notice of Non Payment matching 2026-09-26 is missing"),
            },
        )
        code, payload = self._run(browser)
        self.assertEqual(code, 2)
        self.assertEqual(payload["status"], "HELD")
        self.assertEqual(browser.downloaded, ["860521214", "879512352"])
        folder = self.root / DAY.isoformat()
        saved = list(folder.glob("*.pdf"))
        self.assertEqual([path.name for path in saved], ["860521214 - NOC - Non Payment.pdf"])
        self.assertEqual(saved[0].read_bytes(), PDF_A)
        manifest = json.loads((folder / "manifest.json").read_text())
        self.assertEqual(manifest["status"], "HELD")
        self.assertEqual(manifest["successful_noc"], 1)
        self.assertEqual(manifest["pdfs"], 1)
        self.assertEqual(manifest["policies_on_report"], 2)
        held = [row for row in manifest["policies"] if row["disposition"] == "held"]
        self.assertEqual(held[0]["policy_number"], "879512352")
        self.assertIn("missing", held[0]["reason"])
        readme = (folder / "README.md").read_text()
        self.assertIn("860521214 - NOC - Non Payment.pdf", readme)
        self.assertIn("879512352", readme)

    def test_replay_skips_download_when_the_ledger_matches(self):
        policy = _policy("860521214")
        browser = FakeBrowser(
            ReportCapture(PNG, _report((policy,)), None),
            {"860521214": PDF_A},
        )
        first, payload = self._run(browser)
        self.assertEqual(first, 0)
        self.assertEqual(payload["status"], "PULLED")
        second, again = self._run(browser)
        self.assertEqual(second, 0)
        self.assertEqual(again["status"], "PULLED")
        self.assertEqual(browser.downloaded, ["860521214"])
        manifest = json.loads((self.root / DAY.isoformat() / "manifest.json").read_text())
        self.assertEqual(manifest["policies"][0]["disposition"], "already_present")

    def test_conflicting_pdf_holds_without_a_second_download(self):
        policy = _policy("860521214")
        browser = FakeBrowser(
            ReportCapture(PNG, _report((policy, _policy("879512352"))), None),
            {"860521214": PDF_A, "879512352": PDF_B},
            root=self.root,
        )
        code, payload = self._run(browser)
        self.assertEqual(code, 2)
        self.assertEqual(payload["status"], "HELD")
        self.assertEqual(browser.downloaded, [])
        self.assertIn("conflicts", payload["reason"])

    def test_screenshot_failure_and_bad_png_do_not_claim_empty(self):
        failed = FakeBrowser(IntakeHold("Businessowner/Contractor GL did not open a single new window"))
        code, payload = self._run(failed)
        self.assertEqual(code, 2)
        self.assertNotIn("output", payload)
        self.assertFalse((self.root / DAY.isoformat() / "manifest.json").exists())
        bad = FakeBrowser(ReportCapture(b"not-a-png", None, None))
        code, payload = self._run(bad)
        self.assertEqual(code, 2)
        self.assertIn("screenshot", payload["reason"])
        self.assertFalse((self.root / DAY.isoformat() / "manifest.json").exists())

    def test_unparsed_report_still_saves_the_screenshot_as_held(self):
        browser = FakeBrowser(ReportCapture(PNG, None, "Pending Cancel report is missing or ambiguous"))
        code, payload = self._run(browser)
        self.assertEqual(code, 2)
        folder = self.root / DAY.isoformat()
        manifest = json.loads((folder / "manifest.json").read_text())
        self.assertEqual(manifest["status"], "HELD")
        self.assertEqual(manifest["pdfs"], 0)
        self.assertTrue((folder / "pending-cancel-nonpayment-2026-09-26.png").exists())
        self.assertIn("ambiguous", payload["reason"])

    def test_upload_drive_fails_closed_without_calling_google(self):
        browser = FakeBrowser(ReportCapture(PNG, _report((), blank=True, source="page"), None))
        code, payload = self._run(browser, "--upload-drive")
        self.assertEqual(code, 2)
        self.assertEqual(payload["status"], "HELD")
        self.assertIn("not available", payload["reason"])
        self.assertIn(DRIVE_QA_PARENT_ID, payload["reason"])
        manifest = json.loads((self.root / DAY.isoformat() / "manifest.json").read_text())
        self.assertEqual(manifest["status"], "EMPTY")
        self.assertEqual(manifest["drive"]["status"], "HELD")
        self.assertIn("Progressive BOP", manifest["drive"]["reason"])
        readme = (self.root / DAY.isoformat() / "README.md").read_text()
        self.assertIn("Drive: HELD", readme)
        self.assertNotIn("googleapiclient", module_import_names())

    def test_production_and_unset_env_refuse_before_output(self):
        browser = FakeBrowser(ReportCapture(PNG, _report((), blank=True), None))
        for env in ("PRODUCTION", ""):
            target = Path(self.tmp.name) / f"env-{env or 'unset'}"
            stdout = io.StringIO()
            with patch.dict(os.environ, {"ROBIE_ENV": env}), patch("sys.stdout", stdout):
                code = main(
                    ["--report-date", DAY.isoformat(), "--output-root", str(target)],
                    browser_factory=lambda args: browser,
                )
            self.assertEqual(code, 2, env)
            self.assertFalse(target.exists(), env)
            self.assertIn("HELD", stdout.getvalue())


COMMUNICATIONS_URL = (
    "https://www.foragentsonly.com/managepolicies/policyactivity/"
    "processeddateresults/underwritinglegacy/"
)


class _Query:
    def __init__(self, page, kind):
        self.page = page
        self.kind = kind

    def count(self):
        return self.page.query_count(self.kind)

    def is_visible(self):
        return self.page.query_visible(self.kind)

    def click(self):
        self.page.query_click(self.kind)

    def or_(self, other):
        return _Combo(self.page, "or", self, other)

    def and_(self, other):
        return _Combo(self.page, "and", self, other)


class _Combo:
    def __init__(self, page, mode, left, right):
        self.page = page
        self.mode = mode
        self.left = left
        self.right = right

    def count(self):
        if self.mode == "or":
            return self.left.count() + self.right.count()
        return self.page.and_count(self.left, self.right)

    def is_visible(self):
        if self.count() != 1:
            return False
        if self.mode == "or":
            return any(part.count() == 1 and part.is_visible() for part in (self.left, self.right))
        return self.left.is_visible() and self.right.is_visible()

    def click(self):
        chosen = self.left if self.left.count() else self.right
        chosen.click()

    def or_(self, other):
        return _Combo(self.page, "or", self, other)

    def and_(self, other):
        return _Combo(self.page, "and", self, other)


class _AttachedShell(FakePage):
    """FAO tab that can sit on Communications until Manage Policies Home is clicked."""

    def __init__(
        self,
        url,
        *,
        body="Streetsmart Risk Mgr (33617)",
        home="unique",
        main_nav="one",
        lands=True,
        land_url="https://www.foragentsonly.com/managepolicies/",
    ):
        report = FakePage(
            {("link", "View Reports"), ("button", "Pending Cancel for Nonpayment")},
            url="https://www.foragentsonly.com/bop",
        )
        super().__init__(
            {("link", "Manage Policies"), ("link", "Businessowner/Contractor GL")},
            body=body,
            url=url,
            popup=report,
        )
        self.home = home
        self.main_nav = main_nav
        self.lands = lands
        self.land_url = land_url

    def locator(self, selector):
        if selector == MANAGE_POLICIES_CSS:
            return _Query(self, "css")
        return super().locator(selector)

    def get_by_role(self, role, name=None, exact=True):
        if name == "Main Navigation":
            return _Query(self, f"main-{role}")
        if isinstance(name, re.Pattern) and name.pattern == r"^Manage Policies":
            return _Query(self, f"role-{role}")
        return super().get_by_role(role, name=name, exact=exact)

    def query_count(self, kind):
        if kind == "css":
            if self.home in {"missing", "role-only"}:
                return 0
            if self.home == "ambiguous":
                return 2
            return 1
        if kind == "role-link":
            if self.home in {"missing", "css-only"}:
                return 0
            if self.home == "ambiguous":
                return 2
            return 1
        if kind in {"role-button", "main-link"}:
            return 0
        if kind == "main-button":
            if self.main_nav == "missing":
                return 0
            if self.main_nav == "ambiguous":
                return 2
            return 1
        raise AssertionError(kind)

    def query_visible(self, kind):
        if self.query_count(kind) != 1:
            return False
        if kind.startswith("main-"):
            return True
        return self.home not in {"hidden", "stuck-hidden"}

    def and_count(self, left, right):
        if self.home == "split":
            return 0
        if left.count() == 1 and right.count() == 1:
            return 1
        return 0

    def query_click(self, kind):
        if kind.startswith("main-"):
            self.clicked.append(("button", "Main Navigation"))
            if self.home == "hidden":
                self.home = "unique"
            return
        self.clicked.append(("home", "Manage Policies Home"))
        if self.lands:
            self.url = self.land_url


class _HomeGuard(FakePage):
    def locator(self, selector):
        if selector == MANAGE_POLICIES_CSS:
            raise AssertionError("home control queried on FAO Home")
        return super().locator(selector)


class AgentContextAndHomeTests(unittest.TestCase):
    def test_helper_is_shared_with_fao_memo(self):
        self.assertIs(assert_agent_context, shared_assert_agent_context)
        self.assertIs(assert_agent_context, fao_assert_agent_context)
        self.assertFalse(any("gemini" in name for name in module_import_names()))

    def test_agent_code_shapes_accept_streetsmart_and_reject_other_agencies(self):
        accepted = (
            "StreetSmart Risk Mgr CA33617",
            "Streetsmart Risk Mgr (33617)",
            "33617",
            "Welcome, Carlo Ferrara\n33617c",
            "33617C",
            "CA33617 (33617) 33617c",
            "Streetsmart Risk Mgr (33617)\nTampa FL 90210",
        )
        for body in accepted:
            with self.subTest(body=body):
                self.assertEqual(agent_codes_in_text(body), frozenset({"CA33617"}))
                assert_agent_context(FakePage(set(), body=body), "ca33617")
        self.assertEqual(agent_codes_in_text("(11111)"), frozenset({"CA11111"}))
        self.assertEqual(agent_codes_in_text("99999c"), frozenset({"CA99999"}))
        self.assertEqual(agent_codes_in_text("Tampa FL 90210"), frozenset())
        self.assertEqual(
            agent_codes_in_text("(33617) (11111)"),
            frozenset({"CA33617", "CA11111"}),
        )
        rejected = (
            "",
            "Welcome",
            "CA33617 CA11111",
            "Streetsmart Risk Mgr (33617) CA11111",
            "(33617) (11111)",
            "33617c 99999c",
            "CA33617 (11111)",
        )
        for body in rejected:
            with self.subTest(body=body):
                with self.assertRaisesRegex(IntakeHold, "agent context is missing or ambiguous"):
                    assert_agent_context(FakePage(set(), body=body), "CA33617")
        with self.assertRaisesRegex(IntakeHold, "agent context is missing or ambiguous"):
            assert_agent_context(FakePage(set(), body="Streetsmart Risk Mgr (33617) 33617c"), "CA11111")

    def test_shell_home_does_not_click_the_home_control(self):
        report = FakePage(
            {("link", "View Reports"), ("button", "Pending Cancel for Nonpayment")},
            url="https://www.foragentsonly.com/bop",
        )
        urls = (
            "https://www.foragentsonly.com/",
            "https://www.foragentsonly.com/home",
            "https://www.foragentsonly.com/managepolicies/",
            "https://www.foragentsonly.com/managepolicies/home",
            "https://www.foragentsonly.com/landingpages/managepolicies",
            "https://www.foragentsonly.com/landingpages/managepolicies/",
            "https://www.foragentsonly.com/landingpages/managepolicies/home",
            "https://www.foragentsonly.com/landingpages/managepolicies/home/",
            "https://user:secret@www.foragentsonly.com/home?token=sekret#frag",
            "https://user:secret@www.foragentsonly.com/landingpages/managepolicies/?token=sekret#frag",
        )
        for url in urls:
            with self.subTest(url=url):
                page = _HomeGuard(
                    {("link", "Manage Policies"), ("link", "Businessowner/Contractor GL")},
                    body="Streetsmart Risk Mgr (33617)",
                    url=url,
                    popup=report,
                )
                opened = navigate_to_pending_cancel(page, "CA33617")
                self.assertIs(opened, report)
                already_landing = "/landingpages/managepolicies" in url
                expected = [("link", "Businessowner/Contractor GL")]
                if not already_landing:
                    expected.insert(0, ("goto", "https://www.foragentsonly.com/landingpages/managepolicies/"))
                self.assertEqual(page.clicked, expected)
                self.assertNotIn(("link", "Manage Policies"), page.clicked)
                self.assertNotIn(("home", "Manage Policies Home"), page.clicked)

    def test_communications_tab_opens_home_before_the_agent_assert(self):
        page = _AttachedShell(COMMUNICATIONS_URL, body="Streetsmart Risk Mgr (33617) 33617c")
        opened = navigate_to_pending_cancel(page, "CA33617")
        self.assertIs(opened, page.popup)
        self.assertEqual(
            page.clicked,
            [
                ("home", "Manage Policies Home"),
                ("goto", "https://www.foragentsonly.com/landingpages/managepolicies/"),
                ("link", "Businessowner/Contractor GL"),
            ],
        )
        self.assertEqual(page.url, "https://www.foragentsonly.com/landingpages/managepolicies/")
        self.assertEqual(
            page.popup.clicked,
            [("link", "View Reports"), ("button", "Pending Cancel for Nonpayment")],
        )
        for home in ("css-only", "role-only"):
            with self.subTest(home=home):
                landed = _AttachedShell(COMMUNICATIONS_URL, home=home)
                navigate_to_pending_cancel(landed, "CA33617")
                self.assertEqual(landed.clicked[0], ("home", "Manage Policies Home"))
                self.assertEqual(landed.url, "https://www.foragentsonly.com/landingpages/managepolicies/")

    def test_home_click_that_lands_on_manage_policies_landing_succeeds(self):
        landings = (
            "https://www.foragentsonly.com/landingpages/managepolicies/",
            "https://www.foragentsonly.com/landingpages/managepolicies",
            "https://www.foragentsonly.com/landingpages/managepolicies/home",
            "https://www.foragentsonly.com/landingpages/managepolicies/home/",
        )
        for landing in landings:
            with self.subTest(landing=landing):
                page = _AttachedShell(COMMUNICATIONS_URL, land_url=landing)
                opened = navigate_to_pending_cancel(page, "CA33617")
                self.assertIs(opened, page.popup)
                self.assertEqual(
                    page.clicked,
                    [
                        ("home", "Manage Policies Home"),
                        ("link", "Businessowner/Contractor GL"),
                    ],
                )
                self.assertEqual(page.url, landing)
                self.assertNotIn("gemini", " ".join(str(item) for item in page.clicked))

    def test_near_miss_manage_policies_landing_still_holds(self):
        near = (
            "https://www.foragentsonly.com/landingpages/",
            "https://www.foragentsonly.com/landingpages/managepolicies/policyactivity/",
            "https://www.foragentsonly.com/landingpages/managepolicies/home/extra",
            "http://www.foragentsonly.com/landingpages/managepolicies/",
            "https://evil.example/landingpages/managepolicies/",
        )
        for landing in near:
            with self.subTest(landing=landing):
                page = _AttachedShell(COMMUNICATIONS_URL, land_url=landing)
                with self.assertRaisesRegex(IntakeHold, "FAO Home did not open"):
                    navigate_to_pending_cancel(page, "CA33617")
                self.assertEqual(page.clicked[0], ("home", "Manage Policies Home"))
                self.assertNotIn(("link", "Manage Policies"), page.clicked)
                self.assertNotIn(("link", "Businessowner/Contractor GL"), page.clicked)
                self.assertNotIn("gemini", str(page.clicked).lower())

    def test_shell_home_opens_the_go_to_entry_without_manage_policies(self):
        report = FakePage(
            {("link", "View Reports"), ("button", "Pending Cancel for Nonpayment")},
            url="https://www.foragentsonly.com/bop",
        )
        page = _HomeGuard(
            {("link", "Go to Businessowner/Contractor GL policy search"), ("link", "Manage Policies")},
            body="Streetsmart Risk Mgr (33617)",
            url="https://www.foragentsonly.com/landingpages/managepolicies/",
            popup=report,
        )
        opened = navigate_to_pending_cancel(page, "CA33617")
        self.assertIs(opened, report)
        self.assertEqual(
            page.clicked,
            [("link", "Go to Businessowner/Contractor GL policy search")],
        )
        self.assertEqual(
            report.clicked,
            [("link", "View Reports"), ("button", "Pending Cancel for Nonpayment")],
        )
        button = _HomeGuard(
            {("button", "Go to Businessowner/Contractor GL policy search")},
            body="Streetsmart Risk Mgr (33617)",
            url="https://www.foragentsonly.com/landingpages/managepolicies/home",
            popup=report,
        )
        navigate_to_pending_cancel(button, "CA33617")
        self.assertEqual(
            button.clicked,
            [("button", "Go to Businessowner/Contractor GL policy search")],
        )

    def test_shell_home_holds_when_gl_entry_is_missing_or_ambiguous(self):
        report = FakePage(
            {("link", "View Reports"), ("button", "Pending Cancel for Nonpayment")},
            url="https://www.foragentsonly.com/bop",
        )
        go_to = "Go to Businessowner/Contractor GL policy search"
        cases = (
            set(),
            {("link", "Manage Policies")},
            {("link", "Businessowner/Contractor GL"), ("link", go_to)},
            {("link", "Businessowner/Contractor GL"), ("button", "Businessowner/Contractor GL")},
            {("link", go_to), ("button", go_to)},
        )
        for roles in cases:
            with self.subTest(roles=sorted(roles)):
                page = _HomeGuard(
                    roles,
                    body="Streetsmart Risk Mgr (33617)",
                    url="https://www.foragentsonly.com/landingpages/managepolicies/",
                    popup=report,
                )
                with self.assertRaisesRegex(
                    IntakeHold,
                    "Progressive control 'Businessowner/Contractor GL' is missing or ambiguous",
                ) as caught:
                    navigate_to_pending_cancel(page, "CA33617")
                self.assertEqual(page.clicked, [])
                self.assertNotIn("gemini", str(caught.exception).lower())

    def test_hidden_home_expands_main_navigation_once_without_gemini(self):
        page = _AttachedShell(COMMUNICATIONS_URL, home="hidden")
        navigate_to_pending_cancel(page, "CA33617")
        self.assertEqual(page.clicked[0], ("button", "Main Navigation"))
        self.assertEqual(page.clicked[1], ("home", "Manage Policies Home"))
        self.assertIn(("link", "Businessowner/Contractor GL"), page.clicked)
        self.assertNotIn("gemini", " ".join(str(item) for item in page.clicked))

    def test_missing_or_ambiguous_home_holds_with_the_scrubbed_url(self):
        secret_url = (
            "https://user:secret@www.foragentsonly.com/managepolicies/policyactivity/"
            "processeddateresults/underwritinglegacy/?token=sekret#frag"
        )
        cases = (
            ("missing", "one", True, "Home control was not found"),
            ("ambiguous", "one", True, "Home control was not clicked"),
            ("split", "one", True, "different elements"),
            ("unique", "one", False, "FAO Home did not open"),
            ("stuck-hidden", "one", True, "stayed missing or hidden"),
            ("hidden", "missing", True, "Main Navigation is missing or ambiguous"),
        )
        for home, main_nav, lands, detail in cases:
            with self.subTest(home=home, main_nav=main_nav):
                page = _AttachedShell(secret_url, home=home, main_nav=main_nav, lands=lands)
                with self.assertRaises(IntakeHold) as caught:
                    navigate_to_pending_cancel(page, "CA33617")
                reason = str(caught.exception)
                self.assertIn(detail, reason)
                self.assertIn("tab underwritinglegacy", reason)
                self.assertIn(
                    "page https://www.foragentsonly.com/managepolicies/policyactivity/"
                    "processeddateresults/underwritinglegacy/",
                    reason,
                )
                self.assertNotIn("secret", reason)
                self.assertNotIn("sekret", reason)
                self.assertNotIn("gemini", reason.lower())
                self.assertNotIn(("link", "Businessowner/Contractor GL"), page.clicked)
                self.assertNotIn(("link", "Manage Policies"), page.clicked)



class BopLivePageFixTests(unittest.TestCase):
    def test_partner_sign_on_retry_clicks_once_only_on_popup_blocker_text(self):
        from robie_job_engine.progressive_bop import _retry_partner_sign_on

        blocked = FakePage({("button", "Service Homeowners Policies")}, body="Please disable your pop-up blocker")
        self.assertTrue(_retry_partner_sign_on(blocked))
        self.assertEqual(blocked.clicked, [("button", "Service Homeowners Policies")])
        quiet = FakePage({("button", "Service Homeowners Policies")}, body="Loading")
        self.assertFalse(_retry_partner_sign_on(quiet))
        self.assertEqual(quiet.clicked, [])
        missing = FakePage(set(), body="popup blocker")
        self.assertFalse(_retry_partner_sign_on(missing))

    def test_landing_goto_is_skipped_when_already_there(self):
        from robie_job_engine.progressive_bop import ensure_manage_policies_landing

        page = FakePage(set(), url="https://www.foragentsonly.com/landingpages/managepolicies/")
        ensure_manage_policies_landing(page)
        self.assertEqual(page.clicked, [])
        away = FakePage(set(), url="https://www.foragentsonly.com/home")
        ensure_manage_policies_landing(away)
        self.assertEqual(away.clicked, [("goto", "https://www.foragentsonly.com/landingpages/managepolicies/")])


FIXTURES = Path(__file__).resolve().parents[0] / "fixtures" / "progressive_bop"


def _fixture(name: str) -> bytes:
    return (FIXTURES / name).read_bytes()


def _button_labels(page_html: str) -> list[str]:
    labels = re.findall(r"<button[^>]*>(.*?)</button>", page_html, flags=re.IGNORECASE | re.DOTALL)
    return [_norm_space(label) for label in labels]


def _norm_space(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


class _SilentContext:
    def on(self, event, fn):
        return None

    def remove_listener(self, event, fn):
        return None


class _FinishedDownload:
    def __init__(self, payload: bytes):
        self.payload = payload
        self.finished = False

    def failure(self):
        self.finished = True
        return None

    def save_as(self, path):
        if not self.finished:
            raise AssertionError("save_as ran before the download finished")
        Path(path).write_bytes(self.payload)


class _MissingControl:
    def count(self):
        return 0

    def is_visible(self):
        return False

    def locator(self, selector):
        return self

    def get_by_role(self, role, name=None, exact=True):
        return self

    def wait_for(self, state="visible", timeout=None):
        raise TimeoutError("Timeout 20000ms exceeded")


class _MenuInput:
    def __init__(self, page: "_ExportPage", css: str):
        self.page = page
        self.css = css
        self.value = ""

    def count(self):
        if not self.page.has_inputs:
            return 0
        if self.css == ".report-start":
            return self.page.start_count
        return 1

    def is_visible(self):
        return self.page.menu_open and self.count() == 1

    def wait_for(self, state="visible", timeout=None):
        if state == "visible" and self.is_visible():
            return None
        raise TimeoutError("Timeout 20000ms exceeded")

    def get_attribute(self, name):
        if name == "type":
            return self.page.input_type
        return None

    def fill(self, value):
        if not self.is_visible():
            raise RuntimeError("hidden")
        if not self.page.dates_stick:
            self.value = ""
        else:
            self.value = value
        self.page.date_values[self.css] = self.value

    def input_value(self):
        if self.page.applied and self.page.dates in {"button-text", "unchanged"}:
            return ""
        return self.value


class _ApplyButton:
    def __init__(self, page: "_ExportPage"):
        self.page = page

    def count(self):
        return self.page.apply_count

    def is_visible(self):
        return self.page.menu_open and self.count() == 1

    def wait_for(self, state="visible", timeout=None):
        if state == "visible" and self.is_visible():
            return None
        raise TimeoutError("Timeout 20000ms exceeded")

    def click(self):
        self.page.date_clicks.append("apply")
        self.page.applied = True
        if self.page.dates == "button-text":
            start = self.page.menu.start.value
            end = self.page.menu.end.value
            self.page.button_text = f"{start} - {end}"


class _DateMenu:
    def __init__(self, page: "_ExportPage"):
        self.page = page
        self.start = _MenuInput(page, ".report-start")
        self.end = _MenuInput(page, ".report-end")

    def count(self):
        return 1 if self.page.has_menu else 0

    def is_visible(self):
        return self.page.menu_open

    def locator(self, selector):
        self.page.menu_queries.append(selector)
        if selector == ".report-start":
            return self.start
        if selector == ".report-end":
            return self.end
        return _MissingControl()

    def get_by_role(self, role, name=None, exact=True):
        self.page.menu_roles.append((role, name))
        if role == "button" and name == "Apply" and self.page.apply_count:
            return _ApplyButton(self.page)
        return _MissingControl()


class _DateToggle:
    def __init__(self, page: "_ExportPage", via: str):
        self.page = page
        self.via = via

    def count(self):
        if self.via == "role":
            return self.page.role_count
        return self.page.id_count

    def is_visible(self):
        return self.count() == 1

    def wait_for(self, state="visible", timeout=None):
        if state == "visible" and self.is_visible():
            return None
        raise TimeoutError("Timeout 20000ms exceeded")

    def click(self):
        self.page.date_clicks.append(self.via)
        if self.page.dates != "stuck":
            self.page.menu_open = True

    def inner_text(self):
        return self.page.button_text

    def locator(self, selector):
        self.page.menu_queries.append(selector)
        if (
            self.page.menu_via == "sibling"
            and self.page.has_menu
            and ("dropdown-menu" in selector or "role='menu'" in selector)
        ):
            return self.page.menu
        return _MissingControl()

    def get_attribute(self, name):
        if name == "id":
            return "dropdownMenu2"
        return None


class _ExportPage(FakePage):
    """BOP reports page whose export buttons download canned bytes.

    ``dates`` is ``custom`` (text inputs in the Select Date Range menu),
    ``date`` (type=date), ``missing``, ``ambiguous``, ``id``, ``duplicate-id``,
    ``stuck``, ``no-inputs``, ``reject``, ``no-apply``, ``button-text``,
    ``unchanged``, or ``bad-type``.
    """

    def __init__(
        self,
        files: dict[str, bytes],
        *,
        body: str = "Pending Cancel for Non-Payment",
        dates: str = "custom",
        refresh: str | None = None,
    ):
        super().__init__(
            {("button", name) for name in files},
            body=body,
            url="https://bop.americanstrategic.com/reports",
        )
        self.files = files
        self.context = _SilentContext()
        self.dates = dates
        self.refresh = refresh
        self.dates_stick = dates != "reject"
        self.input_type = "date" if dates == "date" else "number" if dates == "bad-type" else "text"
        self.role_count = 0 if dates in {"missing", "id", "duplicate-id"} else 2 if dates == "ambiguous" else 1
        self.id_count = 0 if dates == "missing" else 2 if dates == "duplicate-id" else 1
        self.has_menu = dates not in {"missing", "no-menu"}
        self.menu_via = "aria" if dates == "aria" else "sibling"
        self.has_inputs = dates != "no-inputs"
        self.start_count = 2 if dates == "duplicate-start" else 1
        self.apply_count = 0 if dates == "no-apply" else 2 if dates == "duplicate-apply" else 1
        self.menu_open = False
        self.applied = False
        self.button_text = "Select Date Range"
        self.date_values: dict[str, str] = {}
        self.date_clicks: list[str] = []
        self.locators: list[str] = []
        self.menu_queries: list[str] = []
        self.menu_roles: list[tuple] = []
        self.role_lookups: list[tuple] = []
        self.states: list[str] = []
        self.events: list[str] = []
        self.exported = False
        self.menu = _DateMenu(self)
        if refresh is not None:
            def wait_for_load_state(state, timeout=None, page=self):
                page.states.append(state)
                page.events.append(state)
                if page.refresh == "timeout":
                    raise TimeoutError("Timeout 20000ms exceeded")

            self.wait_for_load_state = wait_for_load_state

    def locator(self, selector):
        self.locators.append(selector)
        if selector == "#dropdownMenu2":
            return _DateToggle(self, "id")
        if selector == "[aria-labelledby='dropdownMenu2']" and self.menu_via == "aria":
            return self.menu
        return super().locator(selector)

    def get_by_role(self, role, name=None, exact=True):
        self.role_lookups.append((role, name))
        if role == "button" and name == "Select Date Range":
            return _DateToggle(self, "role")
        return super().get_by_role(role, name=name, exact=exact)

    def expect_download(self, timeout=None):
        page = self

        class _Download:
            def __enter__(self):
                return self

            def __exit__(self, exc_type, exc, tb):
                page.exported = True
                page.events.append("export")
                name = page.clicked[-1][1]
                self.value = _FinishedDownload(page.files[name])
                return False

        return _Download()


def _minimal_xlsx(rows: list[list[str]]) -> bytes:
    shared: list[str] = []

    def shared_index(value: str) -> int:
        shared.append(value)
        return len(shared) - 1

    sheet_rows = []
    for row_index, row in enumerate(rows, start=1):
        cells = []
        for column_index, value in enumerate(row):
            column = chr(ord("A") + column_index)
            index = shared_index(value)
            cells.append(f'<c r="{column}{row_index}" t="s"><v>{index}</v></c>')
        sheet_rows.append(f'<row r="{row_index}">{"".join(cells)}</row>')
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    shared_xml = "".join(f"<si><t>{item}</t></si>" for item in shared)
    sheet = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        f'<worksheet xmlns="{ns}"><sheetData>{"".join(sheet_rows)}</sheetData></worksheet>'
    )
    shared_part = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        f'<sst xmlns="{ns}" count="{len(shared)}" uniqueCount="{len(shared)}">{shared_xml}</sst>'
    )
    buffer = io.BytesIO()
    with __import__("zipfile").ZipFile(buffer, "w") as zipped:
        zipped.writestr("[Content_Types].xml", "<Types></Types>")
        zipped.writestr("xl/workbook.xml", f'<workbook xmlns="{ns}"></workbook>')
        zipped.writestr("xl/worksheets/sheet1.xml", sheet)
        zipped.writestr("xl/sharedStrings.xml", shared_part)
    return buffer.getvalue()


class PendingCancelExportTests(unittest.TestCase):
    def test_fixture_page_has_export_buttons_and_no_view_reports(self):
        html = _fixture("reports_page.html").decode("utf-8")
        labels = _button_labels(html)
        self.assertEqual(
            labels,
            [
                "Select Date Range",
                "Apply",
                _PENDING_CANCEL_PDF_EXPORT,
                _PENDING_CANCEL_XLS_EXPORT,
            ],
        )
        self.assertNotIn("View Reports", html)
        self.assertNotIn("VIEW REPORTS", html)
        self.assertNotIn("<select", html.lower())
        self.assertNotIn("<label", html.lower())
        self.assertIn(">Report Dates</p>", html)
        self.assertIn('id="dropdownMenu2"', html)
        self.assertIn('class="dropdown-menu"', html)
        self.assertIn('class="report-start"', html)
        self.assertIn('class="report-end"', html)
        self.assertIn('type="text"', html)
        self.assertLess(html.index('id="dropdownMenu2"'), html.index(_PENDING_CANCEL_PDF_EXPORT))

    def test_export_page_does_not_click_view_reports(self):
        html = _fixture("reports_page.html").decode("utf-8")
        labels = _button_labels(html)
        report = FakePage(
            {("button", label) for label in labels},
            url="https://bop.americanstrategic.com/reports",
            body=html,
        )
        shell = FakePage(
            {("link", "Businessowner/Contractor GL")},
            url="https://www.foragentsonly.com/landingpages/managepolicies/",
            popup=report,
        )
        opened = navigate_to_pending_cancel(shell, "CA33617")
        self.assertIs(opened, report)
        self.assertEqual(report.clicked, [])
        self.assertNotIn(("link", "View Reports"), report.clicked)
        self.assertNotIn(("button", "VIEW REPORTS"), report.clicked)

    def test_network_idle_before_controls_and_no_fixed_sleep(self):
        class _Late(FakePage):
            def __init__(self):
                super().__init__(set(), url="https://bop.americanstrategic.com/reports")
                self.states: list[str] = []

            def wait_for_load_state(self, state, timeout=None):
                self.states.append(state)
                if state == "networkidle":
                    self.roles.update({
                        ("button", _PENDING_CANCEL_PDF_EXPORT),
                        ("button", _PENDING_CANCEL_XLS_EXPORT),
                    })

            def wait_for_timeout(self, _ms):
                raise AssertionError("fixed sleep")

        report = _Late()
        shell = FakePage(
            {("link", "Businessowner/Contractor GL")},
            url="https://www.foragentsonly.com/landingpages/managepolicies/",
            popup=report,
        )
        opened = navigate_to_pending_cancel(shell, "CA33617")
        self.assertIs(opened, report)
        self.assertEqual(report.states, ["networkidle"])
        self.assertEqual(report.clicked, [])

    def test_visible_export_button_is_enough_and_does_not_sleep(self):
        class _Locator(FakeLocator):
            def count(self):
                if not self.page.ready:
                    return 0
                return FakeLocator.count(self)

            def or_(self, _other):
                return self

            def wait_for(self, state="visible", timeout=None):
                self.page.waits.append(state)
                self.page.ready = True
                self.page.roles.update({
                    ("button", _PENDING_CANCEL_PDF_EXPORT),
                    ("button", _PENDING_CANCEL_XLS_EXPORT),
                })

        class _Late(FakePage):
            def __init__(self):
                super().__init__(set(), url="https://bop.americanstrategic.com/reports")
                self.ready = False
                self.waits: list[str] = []

            def get_by_role(self, role, name=None, exact=True):
                return _Locator(self, role=role, name=name)

            def wait_for_load_state(self, state, timeout=None):
                raise AssertionError(state)

            def wait_for_timeout(self, _ms):
                raise AssertionError("fixed sleep")

        report = _Late()
        shell = FakePage(
            {("link", "Businessowner/Contractor GL")},
            url="https://www.foragentsonly.com/landingpages/managepolicies/",
            popup=report,
        )
        opened = navigate_to_pending_cancel(shell, "CA33617")
        self.assertIs(opened, report)
        self.assertEqual(report.waits, ["visible"])
        self.assertEqual(report.clicked, [])

    def test_pdf_list_fixture_extracts_insured_policy_and_cancel_date(self):
        text = _fixture("pending_cancel_list.txt").decode("utf-8")
        page = _ExportPage({_PENDING_CANCEL_PDF_EXPORT: _text_pdf(text)})
        report = read_report_from_page(page, DAY)
        self.assertEqual(report.source, "pdf")
        self.assertEqual(
            [(row.policy_number, row.insured_name) for row in report.policies],
            [("860521214", "3JR Contracting LLC"), ("879512352", "ALTI TRANSPORT LLC")],
        )
        self.assertTrue(all(row.report_date == DAY for row in report.policies))
        self.assertEqual(page.clicked, [("button", _PENDING_CANCEL_PDF_EXPORT)])
        self.assertEqual(page.date_clicks, ["role", "apply"])
        self.assertEqual(
            page.date_values,
            {
                ".report-start": DAY.strftime("%m/%d/%Y"),
                ".report-end": DAY.strftime("%m/%d/%Y"),
            },
        )
        self.assertNotIn(".report-start", page.locators)
        self.assertNotIn(".report-end", page.locators)
        self.assertNotIn("#dropdownMenu2", page.locators)
        self.assertNotIn(("button", "Apply"), page.role_lookups)
        self.assertIn(("button", "Apply"), page.menu_roles)
        with self.assertRaises(IntakeHold):
            policies_from_report_text(
                "860521214 3JR Contracting LLC 09/25/2026",
                report_date=DAY,
            )

    def test_zero_byte_download_holds_and_is_not_an_empty_report(self):
        empty = _fixture("empty_export.bin")
        self.assertEqual(empty, b"")
        page = _ExportPage(
            {
                _PENDING_CANCEL_PDF_EXPORT: empty,
                _PENDING_CANCEL_XLS_EXPORT: empty,
            },
            body="Pending Cancel for Non-Payment. No records.",
        )
        with self.assertRaises(IntakeHold) as caught:
            read_report_from_page(page, DAY)
        reason = str(caught.exception)
        self.assertIn("empty file (0 bytes)", reason)
        self.assertIn("not a report with no policies", reason)
        self.assertNotIn("No records", reason)
        with self.assertRaises(IntakeHold) as extracted:
            report_from_extracted(
                tables=[],
                page_text="Pending Cancel for Nonpayment. No records found.",
                excel_bytes=None,
                pdf_bytes=empty,
                report_date=DAY,
            )
        self.assertIn("empty file (0 bytes)", str(extracted.exception))

    def test_truncated_pdf_fixture_holds(self):
        page = _ExportPage({_PENDING_CANCEL_PDF_EXPORT: _fixture("truncated_export.pdf")})
        with self.assertRaises(IntakeHold) as caught:
            read_report_from_page(page, DAY)
        reason = str(caught.exception)
        self.assertIn("truncated", reason)
        self.assertIn("not a report with no policies", reason)

    def test_html_xls_fixture_is_a_list_when_the_pdf_is_empty(self):
        page = _ExportPage({
            _PENDING_CANCEL_PDF_EXPORT: _fixture("empty_export.bin"),
            _PENDING_CANCEL_XLS_EXPORT: _fixture("pending_cancel_list.html"),
        })
        report = read_report_from_page(page, DAY)
        self.assertEqual(report.source, "excel")
        self.assertEqual(
            [row.policy_number for row in report.policies],
            ["860521214", "879512352"],
        )
        self.assertEqual(report.policies[0].insured_name, "3JR Contracting LLC")
        direct = parse_xls_export(_fixture("pending_cancel_list.html"), report_date=DAY)
        self.assertEqual(direct.policies[1].insured_name, "ALTI TRANSPORT LLC")

    def test_xlsx_zip_parses_without_openpyxl(self):
        blob = _minimal_xlsx([
            ["Policy Number", "Named Insured", "Cancel Date"],
            ["860521214", "3JR Contracting LLC", "09/26/2026"],
        ])
        with patch.dict(sys.modules, {"openpyxl": None}):
            report = parse_excel_report(blob, report_date=DAY)
        self.assertEqual(report.source, "excel")
        self.assertEqual(report.policies[0].policy_number, "860521214")
        self.assertEqual(report.policies[0].insured_name, "3JR Contracting LLC")
        classic = b"\xd0\xcf\x11\xe0" + b"\x00" * 600
        with self.assertRaises(IntakeHold) as caught:
            parse_xls_export(classic, report_date=DAY)
        self.assertIn("older Excel workbook", str(caught.exception))
        self.assertIn("not treated as a report with no policies", str(caught.exception))
        csv_report = parse_xls_export(
            b"Policy Number,Named Insured,Cancel Date\n860521214,3JR Contracting LLC,09/26/2026\n",
            report_date=DAY,
        )
        self.assertEqual(csv_report.policies[0].insured_name, "3JR Contracting LLC")

    def test_a_failed_download_holds_before_the_file_is_read(self):
        from robie_job_engine.progressive_bop import _finished_download_bytes

        class _Failed:
            def failure(self):
                return "net::ERR_ABORTED"

            def save_as(self, path):
                raise AssertionError("save_as ran before the download finished")

        with self.assertRaises(IntakeHold) as caught:
            _finished_download_bytes(_Failed(), label=_PENDING_CANCEL_PDF_EXPORT)
        self.assertIn("did not finish", str(caught.exception))
        self.assertIn("not a report with no policies", str(caught.exception))

    def test_complete_no_records_pdf_is_still_a_blank_report(self):
        page = _ExportPage({
            _PENDING_CANCEL_PDF_EXPORT: _text_pdf("Pending Cancel for Nonpayment. No records."),
        })
        report = read_report_from_page(page, DAY)
        self.assertTrue(report.blank)
        self.assertEqual(report.policies, ())
        self.assertEqual(report.source, "pdf")
        self.assertEqual(page.date_values[".report-start"], DAY.strftime("%m/%d/%Y"))
        self.assertEqual(page.date_values[".report-end"], DAY.strftime("%m/%d/%Y"))

    def test_missing_report_dates_holds_before_export(self):
        page = _ExportPage(
            {_PENDING_CANCEL_PDF_EXPORT: _text_pdf("No records.")},
            dates="missing",
        )
        with self.assertRaises(IntakeHold) as caught:
            read_report_from_page(page, DAY)
        reason = str(caught.exception)
        self.assertIn("Report Dates control is missing or ambiguous", reason)
        self.assertIn("export was not downloaded", reason)
        self.assertFalse(page.exported)
        self.assertEqual(page.clicked, [])

    def test_rejected_date_holds_before_export(self):
        page = _ExportPage(
            {_PENDING_CANCEL_PDF_EXPORT: _text_pdf("No records.")},
            dates="reject",
        )
        with self.assertRaises(IntakeHold) as caught:
            read_report_from_page(page, DAY)
        self.assertIn("report start did not accept", str(caught.exception))
        self.assertIn("export was not downloaded", str(caught.exception))
        self.assertEqual(page.date_clicks, ["role"])
        self.assertFalse(page.exported)

    def test_date_input_uses_iso_and_text_input_uses_month_day_year(self):
        typed = _ExportPage(
            {_PENDING_CANCEL_PDF_EXPORT: _text_pdf("Pending Cancel for Nonpayment. No records.")},
            dates="date",
        )
        choice = apply_bop_report_dates(typed, DAY)
        self.assertEqual(choice.kind, "exact")
        self.assertEqual(choice.start, DAY)
        self.assertEqual(choice.end, DAY)
        self.assertEqual(
            typed.date_values,
            {".report-start": DAY.isoformat(), ".report-end": DAY.isoformat()},
        )
        self.assertEqual(typed.date_clicks, ["role", "apply"])

    def test_id_fallback_is_used_only_when_the_button_name_is_absent(self):
        page = _ExportPage(
            {_PENDING_CANCEL_PDF_EXPORT: _text_pdf("Pending Cancel for Nonpayment. No records.")},
            dates="id",
        )
        apply_bop_report_dates(page, DAY)
        self.assertEqual(page.date_clicks, ["id", "apply"])
        self.assertIn("#dropdownMenu2", page.locators)
        labelled = _ExportPage(
            {_PENDING_CANCEL_PDF_EXPORT: _text_pdf("Pending Cancel for Nonpayment. No records.")},
            dates="aria",
        )
        apply_bop_report_dates(labelled, DAY)
        self.assertEqual(labelled.date_clicks, ["role", "apply"])
        self.assertIn("[aria-labelledby='dropdownMenu2']", labelled.locators)
        self.assertIn(".report-start", labelled.menu_queries)
        ambiguous = _ExportPage(
            {_PENDING_CANCEL_PDF_EXPORT: _text_pdf("No records.")},
            dates="ambiguous",
        )
        with self.assertRaises(IntakeHold) as caught:
            read_report_from_page(ambiguous, DAY)
        self.assertIn("missing or ambiguous", str(caught.exception))
        self.assertNotIn("#dropdownMenu2", ambiguous.locators)
        self.assertEqual(ambiguous.date_clicks, [])
        self.assertFalse(ambiguous.exported)
        duplicate = _ExportPage(
            {_PENDING_CANCEL_PDF_EXPORT: _text_pdf("No records.")},
            dates="duplicate-id",
        )
        with self.assertRaises(IntakeHold) as caught:
            read_report_from_page(duplicate, DAY)
        self.assertIn("missing or ambiguous", str(caught.exception))
        self.assertEqual(duplicate.date_clicks, [])
        self.assertFalse(duplicate.exported)

    def test_menu_that_does_not_open_holds_before_export(self):
        for mode in ("stuck", "no-inputs", "no-menu"):
            with self.subTest(mode=mode):
                page = _ExportPage(
                    {_PENDING_CANCEL_PDF_EXPORT: _text_pdf("No records.")},
                    dates=mode,
                )
                with self.assertRaises(IntakeHold) as caught:
                    read_report_from_page(page, DAY)
                reason = str(caught.exception)
                self.assertIn("still on Select Date Range", reason)
                self.assertIn("were not on the page", reason)
                self.assertIn("export was not downloaded", reason)
                self.assertNotIn("apply", page.date_clicks)
                self.assertFalse(page.exported)

    def test_apply_and_unaccepted_range_hold_before_export(self):
        missing = _ExportPage(
            {_PENDING_CANCEL_PDF_EXPORT: _text_pdf("No records.")},
            dates="no-apply",
        )
        with self.assertRaises(IntakeHold) as caught:
            read_report_from_page(missing, DAY)
        self.assertIn("Apply button is missing or ambiguous", str(caught.exception))
        self.assertFalse(missing.exported)
        duplicate = _ExportPage(
            {_PENDING_CANCEL_PDF_EXPORT: _text_pdf("No records.")},
            dates="duplicate-apply",
        )
        with self.assertRaises(IntakeHold) as caught:
            read_report_from_page(duplicate, DAY)
        self.assertIn("Apply button is missing or ambiguous", str(caught.exception))
        self.assertFalse(duplicate.exported)
        bad_type = _ExportPage(
            {_PENDING_CANCEL_PDF_EXPORT: _text_pdf("No records.")},
            dates="bad-type",
        )
        with self.assertRaises(IntakeHold) as caught:
            read_report_from_page(bad_type, DAY)
        self.assertIn("not a date or text field", str(caught.exception))
        self.assertNotIn("apply", bad_type.date_clicks)
        self.assertFalse(bad_type.exported)
        cleared = _ExportPage(
            {_PENDING_CANCEL_PDF_EXPORT: _text_pdf("No records.")},
            dates="unchanged",
        )
        with self.assertRaises(IntakeHold) as caught:
            read_report_from_page(cleared, DAY)
        self.assertIn("did not accept 2026-09-26", str(caught.exception))
        self.assertIn("apply", cleared.date_clicks)
        self.assertFalse(cleared.exported)

    def test_button_text_is_enough_when_apply_clears_the_inputs(self):
        page = _ExportPage(
            {_PENDING_CANCEL_PDF_EXPORT: _text_pdf("Pending Cancel for Nonpayment. No records.")},
            dates="button-text",
        )
        report = read_report_from_page(page, DAY)
        self.assertTrue(report.blank)
        self.assertEqual(page.button_text, f"{DAY.strftime('%m/%d/%Y')} - {DAY.strftime('%m/%d/%Y')}")
        self.assertTrue(page.exported)
        self.assertEqual(page.clicked, [("button", _PENDING_CANCEL_PDF_EXPORT)])

    def test_duplicate_start_input_holds(self):
        page = _ExportPage(
            {_PENDING_CANCEL_PDF_EXPORT: _text_pdf("No records.")},
            dates="duplicate-start",
        )
        with self.assertRaises(IntakeHold) as caught:
            read_report_from_page(page, DAY)
        self.assertIn("missing or ambiguous", str(caught.exception))
        self.assertFalse(page.exported)

    def test_page_must_finish_loading_after_the_date_is_accepted(self):
        page = _ExportPage(
            {_PENDING_CANCEL_PDF_EXPORT: _text_pdf("Pending Cancel for Nonpayment. No records.")},
            refresh="idle",
        )
        report = read_report_from_page(page, DAY)
        self.assertTrue(report.blank)
        self.assertEqual(page.states, ["networkidle"])
        self.assertTrue(page.exported)
        stalled = _ExportPage(
            {_PENDING_CANCEL_PDF_EXPORT: _text_pdf("No records.")},
            refresh="timeout",
        )
        with self.assertRaises(IntakeHold) as caught:
            read_report_from_page(stalled, DAY)
        self.assertIn("did not finish loading after Report Dates was set", str(caught.exception))
        self.assertFalse(stalled.exported)


if __name__ == "__main__":
    unittest.main()
