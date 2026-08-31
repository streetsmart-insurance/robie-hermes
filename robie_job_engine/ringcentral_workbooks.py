"""Fail-closed ingestion for scheduled RingCentral Excel workbooks.

Only metadata and normalized rows are returned.  The original workbook remains
an ephemeral evidence file on the Test runner and is never suitable for Git.
"""

from __future__ import annotations

import hashlib
import json
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence


REPORT_MARKERS = {
    "daily": "ROBIE_DAILY_CALLS",
    "weekly": "ROBIE_WEEKLY_CALLS",
}
DEFAULT_REQUIRED_SHEETS = {
    "daily": ("Calls",),
    "weekly": ("Users", "Queues", "Calls"),
}
REQUIRED_COLUMNS = {
    "Calls": (
        "Session Id", "From Name", "From Number", "To Name", "To Number",
        "Result", "Call Length", "Handle Time", "Call Start Time",
        "Call Direction", "Queue",
    ),
    "Queues": (
        "Name", "Ext", "# Inbound", "# Answered", "# Abandoned",
        "Avg. Handle Time", "# Holds", "# Refused",
    ),
    # The exact Performance Report user metrics may expand, but identity and
    # extension are the minimum fields required to prove current-user coverage.
    "Users": ("Name", "Ext"),
}
MAX_WORKBOOK_BYTES = 25 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 2_000
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_COMPRESSION_RATIO = 200


class RingCentralEvidenceError(ValueError):
    """The workbook evidence is missing, unsafe, ambiguous, or incomplete."""


def _normalized_marker_text(value: str) -> str:
    return re.sub(r"[^A-Z0-9]+", "_", str(value or "").upper()).strip("_")


def classify_report(*values: str) -> str | None:
    """Classify by the explicit subscription label; ambiguous evidence is refused."""
    text = "_".join(_normalized_marker_text(value) for value in values)
    matches = [
        kind for kind, marker in REPORT_MARKERS.items()
        if re.search(rf"(?:^|_){re.escape(marker)}(?:_|$)", text)
    ]
    if len(matches) > 1:
        raise RingCentralEvidenceError("RingCentral attachment contains conflicting report labels")
    return matches[0] if matches else None


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_xlsx_container(path: Path) -> None:
    """Reject oversized, macro-enabled, externally linked, or malformed XLSX."""
    if path.suffix.casefold() != ".xlsx":
        raise RingCentralEvidenceError("RingCentral evidence must use the .xlsx format")
    size = path.stat().st_size
    if size <= 0 or size > MAX_WORKBOOK_BYTES:
        raise RingCentralEvidenceError("RingCentral workbook size is outside the allowed range")
    if path.read_bytes()[:4] != b"PK\x03\x04":
        raise RingCentralEvidenceError("RingCentral workbook is not a valid XLSX container")
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise RingCentralEvidenceError("RingCentral workbook has too many archive members")
            total = sum(item.file_size for item in members)
            compressed = max(1, sum(item.compress_size for item in members))
            if total > MAX_UNCOMPRESSED_BYTES or total / compressed > MAX_COMPRESSION_RATIO:
                raise RingCentralEvidenceError("RingCentral workbook exceeds decompression limits")
            names = {item.filename.casefold() for item in members}
            if any(".." in Path(item.filename).parts or item.filename.startswith(("/", "\\")) for item in members):
                raise RingCentralEvidenceError("RingCentral workbook contains an unsafe archive path")
            if "xl/vbaproject.bin" in names:
                raise RingCentralEvidenceError("macro-enabled RingCentral workbooks are refused")
            if any(name.startswith("xl/externallinks/") for name in names):
                raise RingCentralEvidenceError("externally linked RingCentral workbooks are refused")
    except zipfile.BadZipFile as exc:
        raise RingCentralEvidenceError("RingCentral workbook is not a readable XLSX archive") from exc


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value).strip()


def _sheet_rows(sheet: Any, sheet_name: str) -> tuple[list[dict[str, str]], list[str]]:
    required = REQUIRED_COLUMNS[sheet_name]
    header_row: int | None = None
    headers: list[str] = []
    for number, values in enumerate(sheet.iter_rows(min_row=1, max_row=25, values_only=True), start=1):
        candidate = [_text(value) for value in values]
        if all(column in candidate for column in required):
            header_row = number
            headers = candidate
            break
    if header_row is None:
        raise RingCentralEvidenceError(
            f"{sheet_name} worksheet is missing required columns: {', '.join(required)}"
        )
    if len(set(item for item in headers if item)) != len([item for item in headers if item]):
        raise RingCentralEvidenceError(f"{sheet_name} worksheet contains duplicate column names")
    rows: list[dict[str, str]] = []
    for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        row = {header: _text(value) for header, value in zip(headers, values) if header}
        if any(row.values()):
            required_values = {
                "Users": ("Name", "Ext"),
                "Queues": ("Name", "Ext", "# Inbound", "# Answered", "# Abandoned", "# Refused"),
                "Calls": (
                    "Session Id", "From Number", "To Number", "Result",
                    "Call Start Time", "Call Direction",
                ),
            }[sheet_name]
            blank = [column for column in required_values if not row.get(column)]
            if blank:
                raise RingCentralEvidenceError(
                    f"{sheet_name} worksheet row {len(rows) + header_row + 1} has blank required values: "
                    f"{', '.join(blank)}"
                )
            rows.append(row)
    return rows, headers


def read_workbook(path: Path, *, required_sheets: Sequence[str] = ()) -> dict[str, Any]:
    """Read allowlisted worksheets with formulas disabled and links refused."""
    validate_xlsx_container(path)
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise RingCentralEvidenceError(f"RingCentral workbook could not be opened: {type(exc).__name__}") from exc
    try:
        missing = [name for name in required_sheets if name not in workbook.sheetnames]
        if missing:
            raise RingCentralEvidenceError(
                f"RingCentral workbook is missing required worksheets: {', '.join(missing)}"
            )
        tables: dict[str, Any] = {}
        for name in ("Users", "Queues", "Calls"):
            if name in workbook.sheetnames:
                rows, columns = _sheet_rows(workbook[name], name)
                tables[name] = {"columns": columns, "rows": rows, "row_count": len(rows)}
        filters = []
        if "Filters" in workbook.sheetnames:
            for values in workbook["Filters"].iter_rows(values_only=True):
                filters.extend(_text(value) for value in values if _text(value))
        return {
            "path": str(path.resolve()),
            "sha256": sha256_file(path),
            "sheets": list(workbook.sheetnames),
            "tables": tables,
            "filters": filters,
        }
    finally:
        workbook.close()


def _fold(value: str) -> str:
    return " ".join(str(value or "").casefold().split())


def _coverage_values(workbooks: Sequence[Mapping[str, Any]], sheet: str) -> set[str]:
    return {
        _fold(row.get("Name", ""))
        for workbook in workbooks
        for row in ((workbook.get("tables") or {}).get(sheet) or {}).get("rows", [])
        if _fold(row.get("Name", ""))
    }


def _filter_coverage(workbooks: Sequence[Mapping[str, Any]], expected: Iterable[str]) -> set[str]:
    cells = [
        _fold(value) for workbook in workbooks for value in workbook.get("filters", [])
    ]
    covered = set()
    for value in expected:
        folded = _fold(value)
        if folded and any(re.search(rf"(?<!\w){re.escape(folded)}(?!\w)", cell) for cell in cells):
            covered.add(folded)
    return covered


def validate_coverage(
    workbooks: Sequence[Mapping[str, Any]],
    *,
    required_users: Sequence[str],
    required_queues: Sequence[str],
) -> dict[str, list[str]]:
    """Prove all configured current users and queues are represented or selected."""
    if not required_users or not required_queues:
        raise RingCentralEvidenceError("current RingCentral users and queues must be explicitly configured")
    actual_users = _coverage_values(workbooks, "Users") or _filter_coverage(workbooks, required_users)
    actual_queues = _coverage_values(workbooks, "Queues") or _filter_coverage(workbooks, required_queues)
    missing_users = [name for name in required_users if _fold(name) not in actual_users]
    missing_queues = [name for name in required_queues if _fold(name) not in actual_queues]
    if missing_users or missing_queues:
        details = []
        if missing_users:
            details.append(f"missing current users: {', '.join(missing_users)}")
        if missing_queues:
            details.append(f"missing current queues: {', '.join(missing_queues)}")
        raise RingCentralEvidenceError("; ".join(details))
    return {"users": list(required_users), "queues": list(required_queues)}


def validate_queue_membership(
    workbooks: Sequence[Mapping[str, Any]],
    *,
    required_queues: Sequence[str],
    required_users: Sequence[str],
    required_queue_members: Mapping[str, Sequence[str]],
    require_observed_legs: bool,
) -> dict[str, list[str]]:
    """Validate the approved queue roster and, weekly, observed member legs."""
    configured = {_fold(queue): [str(member).strip() for member in members] for queue, members in required_queue_members.items()}
    missing_rosters = [queue for queue in required_queues if _fold(queue) not in configured]
    if missing_rosters:
        raise RingCentralEvidenceError(
            f"queue membership is not configured for: {', '.join(missing_rosters)}"
        )
    known_users = {_fold(user) for user in required_users}
    unknown = [
        f"{queue}: {member}"
        for queue in required_queues
        for member in configured.get(_fold(queue), [])
        if _fold(member) not in known_users
    ]
    if unknown:
        raise RingCentralEvidenceError(
            f"queue membership contains users outside the current-user registry: {', '.join(unknown)}"
        )
    if require_observed_legs:
        inbound_by_queue = {
            _fold(row.get("Name", "")): int(str(row.get("# Inbound") or "0").replace(",", ""))
            for workbook in workbooks
            for row in ((workbook.get("tables") or {}).get("Queues") or {}).get("rows", [])
        }
        observed: dict[str, set[str]] = {}
        for workbook in workbooks:
            for row in ((workbook.get("tables") or {}).get("Calls") or {}).get("rows", []):
                if not _fold(row.get("Call Direction", "")).startswith("in") or not _fold(row.get("Queue", "")):
                    continue
                observed.setdefault(_fold(row.get("Queue", "")), set()).add(_fold(row.get("To Name", "")))
        missing_legs = [
            f"{queue}: {member}"
            for queue in required_queues
            if inbound_by_queue.get(_fold(queue), 0) > 0
            for member in configured.get(_fold(queue), [])
            if _fold(member) not in observed.get(_fold(queue), set())
        ]
        if missing_legs:
            raise RingCentralEvidenceError(
                f"current queue members lack observed weekly call legs: {', '.join(missing_legs)}"
            )
    return {queue: configured.get(_fold(queue), []) for queue in required_queues}


def load_evidence_manifest(
    path: Path,
    *,
    expected_kind: str,
    as_of: datetime | None = None,
    max_age_hours: int = 36,
) -> dict[str, Any]:
    """Read a collector manifest and independently re-check every attachment."""
    data = json.loads(path.read_text(encoding="utf-8"))
    if data.get("report_kind") != expected_kind:
        raise RingCentralEvidenceError(
            f"expected {REPORT_MARKERS[expected_kind]} evidence, got {data.get('report_kind') or 'unclassified'}"
        )
    attachments = list(data.get("attachments") or [])
    if not attachments:
        raise RingCentralEvidenceError("RingCentral evidence manifest has no attachments")
    now = as_of or datetime.now(timezone.utc)
    workbooks = []
    for attachment in attachments:
        try:
            received = datetime.fromisoformat(str(attachment.get("received_at") or "").replace("Z", "+00:00"))
            received = received if received.tzinfo else received.replace(tzinfo=timezone.utc)
        except ValueError as exc:
            raise RingCentralEvidenceError("RingCentral evidence attachment lacks a valid received timestamp") from exc
        age_hours = (now - received).total_seconds() / 3600
        if age_hours < 0 or age_hours > max(1, max_age_hours):
            raise RingCentralEvidenceError("RingCentral evidence attachment is stale or future-dated")
        workbook_path = Path(str(attachment.get("path") or "")).expanduser().resolve()
        if not workbook_path.is_file():
            raise RingCentralEvidenceError("RingCentral evidence attachment is missing")
        if sha256_file(workbook_path) != attachment.get("sha256"):
            raise RingCentralEvidenceError("RingCentral evidence attachment checksum changed")
        workbooks.append(read_workbook(workbook_path))
    required_sheets = tuple(data.get("required_sheets") or DEFAULT_REQUIRED_SHEETS[expected_kind])
    available = {sheet for workbook in workbooks for sheet in workbook.get("tables", {})}
    missing = [sheet for sheet in required_sheets if sheet not in available]
    if missing:
        raise RingCentralEvidenceError(f"RingCentral evidence is missing required worksheets: {', '.join(missing)}")
    coverage = validate_coverage(
        workbooks,
        required_users=list(data.get("required_users") or []),
        required_queues=list(data.get("required_queues") or []),
    )
    membership = validate_queue_membership(
        workbooks,
        required_queues=list(data.get("required_queues") or []),
        required_users=list(data.get("required_users") or []),
        required_queue_members=dict(data.get("required_queue_members") or {}),
        require_observed_legs=expected_kind == "weekly",
    )
    return {**data, "workbooks": workbooks, "coverage": coverage, "queue_membership": membership}


def write_evidence_manifest(
    path: Path,
    *,
    report_kind: str,
    attachments: Sequence[Mapping[str, Any]],
    required_sheets: Sequence[str],
    required_users: Sequence[str],
    required_queues: Sequence[str],
    required_queue_members: Mapping[str, Sequence[str]],
    collected_at: datetime | None = None,
) -> Path:
    payload = {
        "schema_version": 1,
        "report_kind": report_kind,
        "report_label": REPORT_MARKERS[report_kind],
        "collected_at": (collected_at or datetime.now(timezone.utc)).isoformat(),
        "required_sheets": list(required_sheets),
        "required_users": list(required_users),
        "required_queues": list(required_queues),
        "required_queue_members": {str(key): list(value) for key, value in required_queue_members.items()},
        "attachments": [dict(item) for item in attachments],
    }
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    return path
