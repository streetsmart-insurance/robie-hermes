"""Fail-closed Gmail collection for scheduled accountability table exports."""

from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Any, Iterable, Mapping, Sequence


MAX_ATTACHMENT_BYTES = 20_000_000
MAX_ARCHIVE_MEMBERS = 1_000
MAX_UNCOMPRESSED_BYTES = 150_000_000
MAX_ROWS = 250_000
MAX_COLUMNS = 256
IGNORED_WORKSHEETS = {"filters", "parameters", "instructions"}


class ScheduledReportEvidenceError(ValueError):
    """The scheduled report evidence is unsafe, stale, ambiguous, or incomplete."""


def _normal(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


def _header(payload: Mapping[str, Any], name: str) -> str:
    wanted = name.casefold()
    for item in payload.get("headers", []) or []:
        if str(item.get("name") or "").casefold() == wanted:
            return str(item.get("value") or "")
    return ""


def _parts(part: Mapping[str, Any]) -> Iterable[Mapping[str, Any]]:
    for child in part.get("parts", []) or []:
        yield child
        yield from _parts(child)


def _labels(report_config: Mapping[str, Mapping[str, Any]]) -> dict[str, str]:
    result: dict[str, str] = {}
    for key, raw in report_config.items():
        label = str((raw or {}).get("label") or "").strip().upper()
        if not label or not re.fullmatch(r"[A-Z0-9_]+", label):
            raise ScheduledReportEvidenceError(f"report {key} requires an exact A-Z/0-9/underscore label")
        if label in result.values():
            raise ScheduledReportEvidenceError(f"duplicate scheduled-report label: {label}")
        result[str(key)] = label
    return result


def classify_report(subject: str, filename: str, report_config: Mapping[str, Mapping[str, Any]]) -> str | None:
    exact = " ".join(str(subject or "").split())
    exact_matches = [
        key for key, raw in report_config.items()
        if " ".join(str((raw or {}).get("exact_subject") or "").split()) == exact
    ]
    if len(exact_matches) > 1:
        raise ScheduledReportEvidenceError("scheduled attachment contains conflicting report labels")
    if exact_matches:
        return exact_matches[0]
    labels = _labels({
        key: raw for key, raw in report_config.items()
        if str((raw or {}).get("label") or "").strip()
    })
    haystack = f"{subject} {filename}".upper()
    matches = [
        key for key, label in labels.items()
        if re.search(rf"(?<![A-Z0-9_]){re.escape(label)}(?![A-Z0-9_])", haystack)
    ]
    if len(matches) > 1:
        raise ScheduledReportEvidenceError("scheduled attachment contains conflicting report labels")
    return matches[0] if matches else None


def _decode_attachment(service: Any, message_id: str, part: Mapping[str, Any]) -> bytes:
    body = part.get("body", {}) or {}
    attachment_id = body.get("attachmentId")
    if attachment_id:
        response = service.users().messages().attachments().get(
            userId="me", messageId=message_id, id=attachment_id
        ).execute()
        encoded = str(response.get("data") or "")
    else:
        encoded = str(body.get("data") or "")
    if not encoded:
        raise ScheduledReportEvidenceError("scheduled report attachment has no data")
    try:
        content = base64.urlsafe_b64decode(encoded + "===")
    except Exception as exc:
        raise ScheduledReportEvidenceError("scheduled report attachment is not valid base64") from exc
    if not content or len(content) > MAX_ATTACHMENT_BYTES:
        raise ScheduledReportEvidenceError("scheduled report attachment size is outside the allowed range")
    return content


def _validate_required_columns(headers: Sequence[str], required: Sequence[str], report_key: str) -> None:
    normalized = {_normal(value) for value in headers if str(value).strip()}
    missing = [value for value in required if _normal(value) not in normalized]
    if missing:
        raise ScheduledReportEvidenceError(
            f"scheduled report {report_key} is missing required columns: {', '.join(missing)}"
        )


def _csv_rows(content: bytes, *, report_key: str, required_columns: Sequence[str]) -> list[list[Any]]:
    if b"\x00" in content:
        raise ScheduledReportEvidenceError(f"scheduled report {report_key} CSV contains NUL bytes")
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ScheduledReportEvidenceError(f"scheduled report {report_key} CSV is not UTF-8") from exc
    rows: list[list[Any]] = []
    for row in csv.reader(io.StringIO(text)):
        if len(row) > MAX_COLUMNS:
            raise ScheduledReportEvidenceError(f"scheduled report {report_key} has too many columns")
        rows.append(list(row))
        if len(rows) > MAX_ROWS:
            raise ScheduledReportEvidenceError(f"scheduled report {report_key} has too many rows")
    if not rows or not any(str(value).strip() for value in rows[0]):
        raise ScheduledReportEvidenceError(f"scheduled report {report_key} has no header row")
    _validate_required_columns([str(value) for value in rows[0]], required_columns, report_key)
    return rows


def _validate_xlsx_container(content: bytes, report_key: str) -> None:
    if not content.startswith(b"PK"):
        raise ScheduledReportEvidenceError(f"scheduled report {report_key} is not a valid XLSX container")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise ScheduledReportEvidenceError(f"scheduled report {report_key} XLSX has too many members")
            if sum(item.file_size for item in members) > MAX_UNCOMPRESSED_BYTES:
                raise ScheduledReportEvidenceError(f"scheduled report {report_key} XLSX exceeds decompression limits")
            for item in members:
                path = PurePosixPath(item.filename)
                if path.is_absolute() or ".." in path.parts:
                    raise ScheduledReportEvidenceError(f"scheduled report {report_key} XLSX has an unsafe path")
                lower = item.filename.casefold()
                if lower.endswith("vbaproject.bin") or "externallinks/" in lower:
                    raise ScheduledReportEvidenceError(
                        f"scheduled report {report_key} XLSX contains macros or external links"
                    )
    except zipfile.BadZipFile as exc:
        raise ScheduledReportEvidenceError(f"scheduled report {report_key} XLSX is unreadable") from exc


def _xlsx_rows(
    content: bytes, *, report_key: str, worksheet: str, required_columns: Sequence[str]
) -> tuple[list[list[Any]], str]:
    _validate_xlsx_container(content, report_key)
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise ScheduledReportEvidenceError(
            f"scheduled report {report_key} XLSX could not be opened: {type(exc).__name__}"
        ) from exc
    selectable = [sheet for sheet in workbook.sheetnames if sheet.casefold() not in IGNORED_WORKSHEETS]
    if worksheet:
        if worksheet not in workbook.sheetnames:
            raise ScheduledReportEvidenceError(f"scheduled report {report_key} is missing worksheet: {worksheet}")
        selected = worksheet
    elif len(selectable) == 1:
        selected = selectable[0]
    else:
        raise ScheduledReportEvidenceError(
            f"scheduled report {report_key} requires an explicit worksheet; available: {', '.join(selectable)}"
        )
    rows: list[list[Any]] = []
    for raw in workbook[selected].iter_rows(values_only=True):
        row = list(raw)
        while row and row[-1] is None:
            row.pop()
        if len(row) > MAX_COLUMNS:
            raise ScheduledReportEvidenceError(f"scheduled report {report_key} has too many columns")
        rows.append(row)
        if len(rows) > MAX_ROWS:
            raise ScheduledReportEvidenceError(f"scheduled report {report_key} has too many rows")
    if not rows or not any(str(value or "").strip() for value in rows[0]):
        raise ScheduledReportEvidenceError(f"scheduled report {report_key} has no header row")
    _validate_required_columns([str(value or "") for value in rows[0]], required_columns, report_key)
    return rows, selected


def _allowed_sender(sender: str, config: Mapping[str, Any]) -> bool:
    normalized = sender.casefold()
    addresses = {str(value).strip().casefold() for value in config.get("allowed_senders", []) if str(value).strip()}
    domains = {str(value).strip().casefold().lstrip("@") for value in config.get("allowed_sender_domains", []) if str(value).strip()}
    if not addresses and not domains:
        raise ScheduledReportEvidenceError("scheduled report collection requires a sender allowlist")
    match = re.search(r"<([^<>]+)>", normalized)
    address = (match.group(1) if match else normalized).strip()
    domain = address.rsplit("@", 1)[-1] if "@" in address else ""
    return address in addresses or domain in domains


def collect_scheduled_tabular_reports(
    service: Any, *, output_dir: Path, config: Mapping[str, Any], as_of: datetime | None = None
) -> dict[str, Any]:
    reports = {str(key): dict(value or {}) for key, value in dict(config.get("reports") or {}).items()}
    if not reports:
        raise ScheduledReportEvidenceError("scheduled report collection requires report definitions")
    labeled = {key: raw for key, raw in reports.items() if str(raw.get("label") or "").strip()}
    if labeled:
        _labels(labeled)
    now = as_of or datetime.now(timezone.utc)
    max_age_hours = max(1, int(config.get("max_age_hours") or 192))
    output_dir.mkdir(parents=True, exist_ok=True)
    message_refs: list[Mapping[str, Any]] = []
    page_token: str | None = None
    for _ in range(10):
        request: dict[str, Any] = {"userId": "me", "q": "has:attachment newer_than:10d", "maxResults": 100}
        if page_token:
            request["pageToken"] = page_token
        response = service.users().messages().list(**request).execute()
        message_refs.extend(response.get("messages", []) or [])
        page_token = str(response.get("nextPageToken") or "").strip() or None
        if not page_token:
            break
    if page_token:
        raise ScheduledReportEvidenceError("scheduled report mailbox scan exceeded the bounded page limit")
    candidates: dict[str, list[dict[str, Any]]] = {key: [] for key in reports}
    seen: set[str] = set()
    for message_meta in message_refs:
        message_id = str(message_meta.get("id") or "")
        if not message_id:
            continue
        message = service.users().messages().get(userId="me", id=message_id).execute()
        received = datetime.fromtimestamp(int(message.get("internalDate") or 0) / 1000, tz=timezone.utc)
        age_hours = (now - received).total_seconds() / 3600
        if age_hours < 0 or age_hours > max_age_hours:
            continue
        payload = message.get("payload", {}) or {}
        if not _allowed_sender(_header(payload, "From"), config):
            continue
        subject = _header(payload, "Subject")
        for part in _parts(payload):
            filename = str(part.get("filename") or "")
            suffix = Path(filename).suffix.casefold()
            if suffix not in {".csv", ".xlsx"}:
                continue
            report_key = classify_report(subject, filename, reports)
            if report_key is None:
                continue
            content = _decode_attachment(service, message_id, part)
            digest = hashlib.sha256(content).hexdigest()
            if digest in seen:
                continue
            seen.add(digest)
            definition = reports[report_key]
            required_columns = [str(value) for value in definition.get("required_columns", [])]
            worksheet = str(definition.get("worksheet") or "").strip()
            if suffix == ".csv":
                rows = _csv_rows(content, report_key=report_key, required_columns=required_columns)
                selected_sheet = ""
            else:
                rows, selected_sheet = _xlsx_rows(
                    content, report_key=report_key, worksheet=worksheet, required_columns=required_columns
                )
            candidates[report_key].append({
                "received_at": received,
                "sha256": digest,
                "filename_sha256": hashlib.sha256(filename.encode()).hexdigest()[:16],
                "message_id_sha256": hashlib.sha256(message_id.encode()).hexdigest()[:16],
                "rows": rows,
                "worksheet": selected_sheet,
            })
    sources: dict[str, str] = {}
    trackers: dict[str, str] = {}
    evidence: dict[str, Any] = {}
    for report_key, definition in reports.items():
        available = sorted(candidates[report_key], key=lambda item: item["received_at"], reverse=True)
        if not available:
            if bool(definition.get("required", True)):
                raise ScheduledReportEvidenceError(f"no fresh scheduled report found for {report_key}")
            evidence[report_key] = {"source_status": "not supplied", "required": False}
            continue
        selected = available[0]
        target = output_dir / f"{report_key}-{selected['received_at']:%Y%m%dT%H%M%SZ}-{selected['sha256'][:12]}.csv"
        with target.open("w", encoding="utf-8", newline="") as handle:
            csv.writer(handle).writerows(selected["rows"])
        destination = str(definition.get("destination") or "source")
        if destination == "tracker":
            trackers[report_key] = str(target.resolve())
        elif destination == "source":
            sources[report_key] = str(target.resolve())
        else:
            raise ScheduledReportEvidenceError(f"report {report_key} has unsupported destination: {destination}")
        evidence[report_key] = {
            "source_status": "available",
            "received_at": selected["received_at"].isoformat(),
            "attachment_sha256": selected["sha256"],
            "normalized_csv_sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
            "row_count": max(0, len(selected["rows"]) - 1),
            "worksheet": selected["worksheet"],
            "message_id_sha256": selected["message_id_sha256"],
            "filename_sha256": selected["filename_sha256"],
        }
    manifest = output_dir / f"scheduled-reports-{now:%Y%m%dT%H%M%SZ}.json"
    manifest.write_text(json.dumps({
        "collected_at": now.isoformat(), "sources": sources, "trackers": trackers, "evidence": evidence,
    }, indent=2), encoding="utf-8")
    return {"sources": sources, "trackers": trackers, "manifest": str(manifest.resolve()), "evidence": evidence}
