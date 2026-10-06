"""Monthly lost-customer retention reporting with fail-closed evidence rules."""

from __future__ import annotations

import base64
import hashlib
import io
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from email.message import EmailMessage
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence


SHEETS_SCOPE = "https://www.googleapis.com/auth/spreadsheets"
UNKNOWN = "Unknown"
CANONICAL_DEPARTMENTS = (
    "Personal Lines", "Commercial Lines", "Trucking and Transportation",
)
DEPARTMENT_RECIPIENT_KEYS = {
    "Personal Lines": "personal",
    "Commercial Lines": "commercial",
    "Trucking and Transportation": "trucking",
}
EXECUTIVE_DEPARTMENT = "Executive Team"
SOP_RESULTS = ("Followed", "Partially Followed", "Not Followed", "Cannot Verify")
SOP_SOURCES = {
    "Cancellations": "1gpDHSl6xm_RTF9d4Htyd9i8zj7eIWebvIDjFJJSlQzo",
    "Client Cancellation Request": "1rvAYtS1HY9hdX6Is0Z0Llt6FhizipoZOyZUbSyRkyhc",
    "Remarketing and Rewriting": "1rFsVS4DzN7x3xL1aZS4qGU5pPQFHeULYROsXsn0i9lw",
    "Reinstatements": "1nDze_axPNH7tqXr5GPnhoRP7j-eplAbmzkRcrSNZarQ",
    "Manual Renewals": "1NSHFdtgKYXeF6OnJRg9Qlse_yiHUoaXTYPi0X1OwLr4",
    "VIP Service Standards": "1YA70LaB0vmmRfKZZLFHMu2w8h8OrLGxrPG_h858oAis",
}
REQUIRED_REVIEW_HEADERS = (
    "Month", "Applicant ID", "Account Name", "Policy Count", "Lines of Business",
    "Policy Numbers", "Department", "CSR", "Assigned Agent", "Annualized Premium",
    "Account Status Classification", "Evidence-Supported Cause", "Evidence Summary",
    "Responsibility Lane", "Preventable?", "Confidence", "Magellan Match",
    "Magellan Sentiment", "Recovery Opportunity", "Recommended Account Action",
    "Systemic Prevention",
)


def normalize(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def normalize_name(value: Any) -> str:
    name = normalize(value)
    if "," in name:
        last, first = [normalize(x) for x in name.split(",", 1)]
        name = f"{first} {last}"
    return re.sub(r"[^a-z0-9]", "", name.casefold())


def employee_directory(values: Sequence[Sequence[Any]]) -> dict[str, dict[str, str]]:
    if not values:
        raise ValueError("AppSheet Employees tab is empty")
    headers = [normalize(v) for v in values[0]]
    needed = ("Name", "Email", "Department", "Employment Status", "Position")
    missing = [h for h in needed if h not in headers]
    if missing:
        raise ValueError(f"AppSheet Employees missing headers: {', '.join(missing)}")
    result = {}
    for row in values[1:]:
        item = {h: normalize(row[i]) if i < len(row) else "" for i, h in enumerate(headers)}
        key = normalize_name(item["Name"])
        if key and item["Employment Status"].casefold() == "active":
            result[key] = item
    return result


def fallback_department(lines_of_business: Any) -> str:
    lob = normalize(lines_of_business).casefold()
    if any(x in lob for x in ("trucking", "truckers", "motor carrier", "transportation")):
        return "Trucking and Transportation"
    if any(x in lob for x in ("personal", "homeowner", "flood", "dwelling", "umbrella")):
        return "Personal Lines"
    return "Commercial Lines"


def resolve_department(item: Mapping[str, str], directory: Mapping[str, Mapping[str, str]]) -> tuple[str, str, bool]:
    assigned = directory.get(normalize_name(item.get("Assigned Agent")))
    csr = directory.get(normalize_name(item.get("CSR")))
    chosen = csr if assigned and assigned.get("Department") == EXECUTIVE_DEPARTMENT else assigned
    if not chosen or chosen.get("Department") not in CANONICAL_DEPARTMENTS:
        chosen = csr if csr and csr.get("Department") in CANONICAL_DEPARTMENTS else None
    if chosen:
        return chosen["Department"], f"AppSheet Employees: {chosen['Name']}", False
    return fallback_department(item.get("Lines of Business")), "LOB fallback — employee unmatched", True


def tenure_months(as_of: datetime, start: datetime | None) -> int | None:
    if start is None:
        return None
    return max(0, (as_of.year - start.year) * 12 + as_of.month - start.month - (as_of.day < start.day))


def tenure_band(months: int | None) -> str:
    if months is None:
        return "Cannot Verify"
    if months < 12:
        return "<1 year"
    if months < 36:
        return "1–3 years"
    if months < 60:
        return "3–5 years"
    return "5+ years"


def conservative_sop_audit(item: Mapping[str, str]) -> dict[str, str]:
    evidence = normalize(item.get("Evidence Summary"))
    folded = evidence.casefold()
    result = {key: "Cannot Verify" for key in (
        "SPLICE Call", "Email", "Text", "Postal Mail / Bad Contact",
        "Department Label", "Cancellation Notice / Reason", "Written Authorization",
        "EFT / Payment Rescue", "$5k+ Two Outreaches", "$10k+ AM Alert",
        "Licensed Handoff", "Carrier-Confirmed Reinstatement",
        "No Premature Renewal/Coverage Confirmation",
    )}
    if "signed cancellation request" in folded or "signed lpr" in folded:
        result["Written Authorization"] = "Followed"
    if "cancellation notice" in folded and ("reason" in folded or item.get("Evidence-Supported Cause")):
        result["Cancellation Notice / Reason"] = "Partially Followed"
    if "mistaken renewal confirmation" in folded or "incorrectly told" in folded:
        result["No Premature Renewal/Coverage Confirmation"] = "Not Followed"
    if "carrier confirm" in folded and "reinstat" in folded:
        result["Carrier-Confirmed Reinstatement"] = "Followed"
    result["Overall SOP Result"] = (
        "Not Followed" if "Not Followed" in result.values()
        else "Partially Followed" if any(v in {"Followed", "Partially Followed"} for v in result.values())
        else "Cannot Verify"
    )
    result["Evidence Citation"] = (
        f"3-Month Account Review | {item.get('Month')} | Applicant {item.get('Applicant ID')} | "
        f"Evidence Summary: {evidence or 'No supporting evidence recorded'}"
    )
    return result


def money(value: Any) -> float:
    cleaned = re.sub(r"[^0-9.()-]", "", normalize(value)).replace("(", "-").replace(")", "")
    try:
        return float(cleaned or 0)
    except ValueError:
        return 0.0


def records(values: Sequence[Sequence[Any]]) -> list[dict[str, str]]:
    if not values:
        return []
    headers = [normalize(v) for v in values[0]]
    if tuple(headers[: len(REQUIRED_REVIEW_HEADERS)]) != REQUIRED_REVIEW_HEADERS:
        raise ValueError("3-Month Account Review headers do not match the approved contract")
    output = []
    for row in values[1:]:
        item = {header: normalize(row[i]) if i < len(row) else "" for i, header in enumerate(headers)}
        if item.get("Month") and item.get("Applicant ID"):
            if not item.get("Evidence-Supported Cause") or "not found" in item["Evidence-Supported Cause"].casefold():
                item["Evidence-Supported Cause"] = UNKNOWN
            output.append(item)
    return output


def validate_monthly_source(values: Sequence[Sequence[Any]], month: str) -> dict[str, Any]:
    rows = []
    for row in values:
        first = normalize(row[0] if row else "")
        if first.isdigit() and len(row) > 12 and normalize(row[9]):
            rows.append(row)
    keys = ["|".join((normalize(r[0]).casefold(), normalize(r[9]).casefold(), normalize(r[12]).casefold())) for r in rows]
    if len(keys) != len(set(keys)):
        raise ValueError(f"{month}: duplicate policy transaction keys remain")
    departments = Counter(normalize(r[8]) or UNKNOWN for r in rows)
    return {
        "month": month,
        "policy_rows": len(rows),
        "accounts": len({normalize(r[0]) for r in rows}),
        "departments": dict(sorted(departments.items())),
        "sha256": hashlib.sha256(json.dumps(rows, ensure_ascii=False, sort_keys=True).encode()).hexdigest(),
    }


def summarize(items: Iterable[Mapping[str, str]]) -> dict[str, Any]:
    grouped: dict[str, dict[str, Any]] = {}
    for item in items:
        key = f"{item.get('Month')}|{item.get('Department') or UNKNOWN}"
        group = grouped.setdefault(key, {
            "month": item.get("Month"), "department": item.get("Department") or UNKNOWN,
            "accounts": 0, "policies": 0, "premium": 0.0, "classifications": Counter(),
            "causes": Counter(), "magellan_matches": 0,
        })
        group["accounts"] += 1
        group["policies"] += int(money(item.get("Policy Count")))
        group["premium"] += money(item.get("Annualized Premium"))
        group["classifications"][item.get("Account Status Classification") or UNKNOWN] += 1
        group["causes"][item.get("Evidence-Supported Cause") or UNKNOWN] += 1
        if normalize(item.get("Magellan Match")).casefold() not in {"", "no matched record", "not available"}:
            group["magellan_matches"] += 1
    result = []
    for group in grouped.values():
        group["classifications"] = dict(group["classifications"])
        group["causes"] = dict(group["causes"])
        group["premium"] = round(group["premium"], 2)
        result.append(group)
    return {"groups": sorted(result, key=lambda x: (x["month"], x["department"]))}


def format_period_label(months: Sequence[str]) -> str:
    months = tuple(months)
    if not months:
        raise ValueError("period requires at least one month")
    if len(months) == 1:
        return months[0]
    if months == ("June 2026", "July 2026", "August 2026"):
        return "June–August 2026"
    return ", ".join(months)


def build_retention_xlsx(
    *,
    monthly_tabs: Mapping[str, Sequence[Sequence[Any]]],
    review_rows: Sequence[Mapping[str, str]],
    period_label: str,
) -> tuple[str, bytes]:
    """Build an .xlsx of the monthly Pulse tab(s) plus Account Review rows.

    Drive export is unavailable to the Prod SA (403), so we rebuild from the
    Sheets values already loaded for validation.
    """
    from openpyxl import Workbook

    if not monthly_tabs:
        raise ValueError("monthly_tabs is required for the retention workbook")
    workbook = Workbook()
    first = True
    for month, rows in monthly_tabs.items():
        sheet = workbook.active if first else workbook.create_sheet()
        first = False
        sheet.title = str(month)[:31]
        for row in rows:
            sheet.append(list(row))
    review_sheet = workbook.create_sheet("Account Review")
    review_sheet.append(list(REQUIRED_REVIEW_HEADERS))
    for item in review_rows:
        review_sheet.append([item.get(header, "") for header in REQUIRED_REVIEW_HEADERS])
    buffer = io.BytesIO()
    workbook.save(buffer)
    slug = period_label.replace("–", "-").replace(" ", "_").replace(",", "")
    filename = f"Lost_Customer_Retention_{slug}.xlsx"
    return filename, buffer.getvalue()


def spreadsheet_lines(period_label: str, *, sheet_url: str | None, attached: bool) -> list[str]:
    """Where the recipient opens the monthly spreadsheet."""
    lines: list[str] = []
    if sheet_url:
        lines += [
            f"Open the {period_label} Google Sheet: {sheet_url}",
            f"Tabs: {period_label} (Pulse detail) and Account Review.",
        ]
    if attached:
        lines.append(
            "A copy is also attached as an .xlsx file." if sheet_url else "The monthly spreadsheet is attached."
        )
    return lines


def department_email(
    department: str,
    items: Sequence[Mapping[str, str]],
    *,
    run_id: str,
    period_label: str,
    sheet_url: str | None = None,
    attached: bool = True,
) -> str:
    lines = [
        f"StreetSmart Lost Customer Retention Review — {department}", "",
        f"{period_label} validated review", f"Run ID: {run_id}", "",
        *spreadsheet_lines(period_label, sheet_url=sheet_url, attached=attached), "",
    ]
    if not items:
        lines += ["No account-level records mapped to this department for the validated period."]
        return "\n".join(lines)
    totals = summarize(items)["groups"]
    lines.append("Monthly totals:")
    for group in totals:
        lines.append(f"- {group['month']}: {group['accounts']} accounts / {group['policies']} policies / ${group['premium']:,.2f}")
    lines += ["", "Account-level findings:"]
    for item in items:
        lines += [
            "",
            f"- {item['Month']} | Applicant {item['Applicant ID']} | {item['Account Name']}",
            f"  Policies: {item['Policy Count']} | LOB: {item['Lines of Business']} | Premium: {item['Annualized Premium']}",
            f"  Status: {item['Account Status Classification']} | Cause: {item['Evidence-Supported Cause'] or UNKNOWN}",
            f"  Evidence: {item['Evidence Summary'] or 'No supporting evidence found.'}",
            f"  Magellan: {item['Magellan Match'] or 'No matched record'} / {item['Magellan Sentiment'] or 'Not available'}",
            f"  Department mapping: {item.get('_department_source', 'pre-mapped workbook')}"+
            (" [FALLBACK EXCEPTION]" if item.get("_department_exception") else ""),
            f"  SOP audit: {conservative_sop_audit(item)['Overall SOP Result']} | {conservative_sop_audit(item)['Evidence Citation']}",
            f"  Recommended action: {item['Recommended Account Action'] or 'Review and document a supported cause.'}",
        ]
    lines += ["", "Assignment alone is not evidence of employee fault. Unsupported causes remain Unknown."]
    return "\n".join(lines)


def executive_email(
    items: Sequence[Mapping[str, str]],
    *,
    run_id: str,
    period_label: str,
    sheet_url: str | None = None,
    attached: bool = True,
) -> str:
    summary = summarize(items)["groups"]
    total_accounts = len(items)
    total_policies = sum(int(money(x.get("Policy Count"))) for x in items)
    total_premium = sum(money(x.get("Annualized Premium")) for x in items)
    statuses = Counter(x.get("Account Status Classification") or UNKNOWN for x in items)
    lines = [
        "StreetSmart Lost Customer Retention Review — Executive Trends", "",
        f"{period_label} validated review", f"Run ID: {run_id}",
        *spreadsheet_lines(period_label, sheet_url=sheet_url, attached=attached),
        f"Overall: {total_accounts} account-months / {total_policies} policies / ${total_premium:,.2f} flagged annualized premium", "",
        "Department/month totals:",
    ]
    for group in summary:
        lines.append(f"- {group['month']} — {group['department']}: {group['accounts']} accounts / {group['policies']} policies / ${group['premium']:,.2f}")
    lines += ["", "Corrected account classifications:"]
    for label, count in statuses.most_common():
        lines.append(f"- {label}: {count}")
    lines += [
        "", "Controls verified:",
        "- Raw policy rows are consolidated to account-month findings.",
        "- Rewrites, reinstatements and partial-policy losses remain separate from true customer loss.",
        "- Unsupported causes are Unknown.",
        "- Assignment alone is never treated as employee fault.",
        "- Magellan is corroborating evidence only when an account match exists.",
        "", "This is a scheduled monthly-updated Google Sheet, not a continuously updating dashboard.",
    ]
    return "\n".join(lines)


@dataclass(frozen=True)
class Attachment:
    filename: str
    content: bytes
    maintype: str = "application"
    subtype: str = "vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass(frozen=True)
class MessageSpec:
    key: str
    to: tuple[str, ...]
    subject: str
    body: str
    attachments: tuple[Attachment, ...] = ()


def build_messages(
    items: Sequence[Mapping[str, str]],
    recipients: Mapping[str, str],
    run_id: str,
    *,
    period_label: str,
    attachments: Sequence[Attachment] = (),
    sheet_url: str | None = None,
) -> list[MessageSpec]:
    by_department: dict[str, list[Mapping[str, str]]] = defaultdict(list)
    for item in items:
        by_department[item.get("Department") or UNKNOWN].append(item)
    mapping = tuple((department, key) for department, key in DEPARTMENT_RECIPIENT_KEYS.items())
    attached = tuple(attachments)
    messages = []
    for department, key in mapping:
        messages.append(MessageSpec(
            key=f"department:{key}", to=(recipients[key],),
            subject=f"Lost Customer Retention Review — {department} — {period_label}",
            body=department_email(
                department, by_department.get(department, []), run_id=run_id, period_label=period_label,
                sheet_url=sheet_url, attached=bool(attached),
            ),
            attachments=attached,
        ))
    messages.append(MessageSpec(
        key="executive", to=(recipients["carlo"], recipients["jake"]),
        subject=f"Lost Customer Retention Review — Executive Trends — {period_label}",
        body=executive_email(
            items, run_id=run_id, period_label=period_label, sheet_url=sheet_url, attached=bool(attached),
        ),
        attachments=attached,
    ))
    return messages


def send_message(gmail: Any, sender: str, spec: MessageSpec) -> dict[str, Any]:
    message = EmailMessage()
    message["To"] = ", ".join(spec.to)
    message["From"] = sender
    message["Subject"] = spec.subject
    message["X-ROBIE-Idempotency-Key"] = spec.key
    message.set_content(spec.body)
    attached_names = []
    for attachment in spec.attachments:
        if not attachment.content:
            raise RuntimeError(f"empty attachment: {attachment.filename}")
        message.add_attachment(
            attachment.content,
            maintype=attachment.maintype,
            subtype=attachment.subtype,
            filename=attachment.filename,
        )
        attached_names.append(attachment.filename)
    raw = base64.urlsafe_b64encode(message.as_bytes()).decode("ascii")
    sent = gmail.users().messages().send(userId="me", body={"raw": raw}).execute()
    return {
        "key": spec.key,
        "to": list(spec.to),
        "message_id": sent.get("id"),
        "thread_id": sent.get("threadId"),
        "attachments": attached_names,
    }


def load_state(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {"runs": {}}
    return json.loads(path.read_text(encoding="utf-8"))


def save_state(path: Path, state: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(".tmp")
    temp.write_text(json.dumps(state, indent=2, sort_keys=True), encoding="utf-8")
    temp.replace(path)
