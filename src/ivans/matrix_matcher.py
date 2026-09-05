"""IVANS Exchange Download Matrix Matcher.

Cross-references carrier and Line of Business (LOB) against the agency's
active IVANS Exchange Connections export to distinguish automated downloads
from true manual non-download renewals.
"""

import re
import logging
from pathlib import Path
from typing import Optional, Dict
import openpyxl

logger = logging.getLogger("ivans_matcher")

BASE_DIR = Path(__file__).resolve().parent.parent.parent
DEFAULT_IVANS_XLSX = BASE_DIR / "data" / "ivans" / "ivans_export_1.xlsx"


class IvansMatrixMatcher:
    """Evaluates whether a carrier + LOB has an active electronic IVANS download connection."""

    def __init__(self, matrix_path: Optional[Path] = None):
        self.matrix_path = matrix_path or DEFAULT_IVANS_XLSX
        self.ivans_map: Dict[str, Dict[str, bool]] = {}
        self._load_matrix()

    def _load_matrix(self):
        if not self.matrix_path.exists():
            logger.warning(f"IVANS matrix file not found at {self.matrix_path}")
            return

        try:
            wb = openpyxl.load_workbook(str(self.matrix_path), data_only=True)
            ws = wb["Connections"] if "Connections" in wb.sheetnames else wb.active
            headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]

            for row in ws.iter_rows(min_row=3, values_only=True):
                carrier = str(row[0]).strip() if row[0] is not None else ""
                if not carrier or carrier.lower() in ("carriers", "none"):
                    continue

                c_clean = re.sub(r"[^a-z0-9]", "", carrier.lower())
                if c_clean not in self.ivans_map:
                    self.ivans_map[c_clean] = {}

                for col_idx, val in enumerate(row[1:], start=1):
                    if col_idx < len(headers):
                        lob_name = headers[col_idx]
                        is_dl = val in ["DL", "PK", "C"]
                        self.ivans_map[c_clean][lob_name] = (
                            self.ivans_map[c_clean].get(lob_name, False) or is_dl
                        )
            logger.info(f"Loaded IVANS matrix with {len(self.ivans_map)} carrier records from {self.matrix_path.name}")
        except Exception as e:
            logger.error(f"Failed to parse IVANS matrix from {self.matrix_path}: {e}")

    @staticmethod
    def map_ezlynx_lob(ez_lob: str) -> str:
        """Normalizes EZLynx Line of Business string to standard IVANS column header."""
        lob = (ez_lob or "").lower().strip()
        if "auto" in lob:
            if "personal" in lob or "pers" in lob:
                return "PL AUTOMOBILE"
            return "CL AUTOMOBILE"
        if "inland marine" in lob:
            if "personal" in lob or "pers" in lob:
                return "PL INLAND MARINE"
            return "CL INLAND MARINE"
        if any(w in lob for w in ["workers comp", "work comp", "wc", "workers' compensation"]):
            return "WORKERS COMPENSATION"
        if any(w in lob for w in ["business owners", "bop"]):
            return "BUSINESS OWNERS"
        if any(w in lob for w in ["genl liability", "general liability", "gl"]):
            return "GENERAL LIABILITY"
        if any(w in lob for w in ["homeowners", "ho-", "ho3", "ho4", "ho5", "ho6"]):
            return "HOMEOWNERS"
        if "dwelling" in lob:
            return "DWELLING FIRE"
        if "flood" in lob:
            return "FLOOD"
        if "umbrella" in lob:
            if "pers" in lob:
                return "PL UMBRELLA"
            return "CL UMBRELLA"
        if "property" in lob:
            return "PROPERTY"
        if "package" in lob or "pkg" in lob:
            return "CL PACKAGE"
        return lob.upper()

    def is_ivans_downloading(self, carrier: str, lob: str) -> bool:
        """Returns True if the carrier + LOB has an active automated IVANS download."""
        if not self.ivans_map:
            return False

        ez_c_clean = re.sub(r"[^a-z0-9]", "", (carrier or "").lower())
        # State Assigned Risk pools never download via IVANS
        if any(k in ez_c_clean for k in ["njcrib", "assignedrisk", "nycrib", "wcra"]):
            return False

        ivans_lob = self.map_ezlynx_lob(lob)

        # Match carrier against IVANS trading partners
        for ic, lobs in self.ivans_map.items():
            if ic in ez_c_clean or ez_c_clean in ic or (len(ic) > 5 and ic[:6] == ez_c_clean[:6]):
                if lobs.get(ivans_lob, False):
                    return True
        return False


# Global singleton
ivans_matcher = IvansMatrixMatcher()
